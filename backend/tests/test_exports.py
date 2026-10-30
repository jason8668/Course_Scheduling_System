"""M5-1:课表导出。Excel 以读回比对验证内容;HTML(PDF 来源)验中文与版面;
全校总表/批次/RBAC。PDF/PNG 的 worker 渲染于真实环境另行验证。
"""

import io
import zipfile

import pytest
from openpyxl import load_workbook

from app.models.user import Role
from app.services import timetable_export as tex
from tests.conftest import grant_grades, make_user
from tests.dates import SEM_END, SEM_START  # 日期一律由执行当日推算,不硬编
from tests.test_substitutions import _World

PW = "password123"


@pytest.fixture
def w(env):
    client, db = env
    grant_grades(db, make_user(db, "s", PW, roles=[Role.director]))
    client.post("/api/auth/login", json={"username": "s", "password": PW})
    sid = client.post("/api/semesters", json={
        "academic_year": 2026, "term": 1, "template_key": "junior_high",
        "start_date": SEM_START.isoformat(), "end_date": SEM_END.isoformat(),
    }).json()["id"]
    world = _World(client, db, sid)
    world.teacher("王师", ["国文"])
    world.place("王师", "国文", "701", 0)   # 周三第一节
    world.publish()
    return world


def _cells(ws) -> set[str]:
    out: set[str] = set()
    for row in ws.iter_rows(values_only=True):
        for v in row:
            if v:
                out.add(str(v))
    return out


def _xlsx(resp):
    assert resp.status_code == 200, resp.text
    assert "spreadsheetml" in resp.headers["content-type"]
    return load_workbook(io.BytesIO(resp.content))


# ── 内容一致(Excel 读回)─────────────────────────────────────
def test_class_excel_matches_grid(w):
    cid = w.classes["701"]
    r = w.client.get(f"/api/export/timetable{w.q}&view=class&target_id={cid}&fmt=xlsx")
    wb = _xlsx(r)
    ws = wb.active
    cells = _cells(ws)
    assert "701 课表" in " ".join(cells) or any("课表" in c for c in cells)
    # 周三第一节有「国文」与「王师」
    joined = "\n".join(cells)
    assert "国文" in joined
    assert "王师" in joined
    assert "周三" in cells  # 表头星期
    assert "第一节" in cells  # 节次名


def test_teacher_excel_shows_class_not_teacher(w):
    tid = w.teachers["王师"]
    r = w.client.get(f"/api/export/timetable{w.q}&view=teacher&target_id={tid}&fmt=xlsx")
    wb = _xlsx(r)
    joined = "\n".join(_cells(wb.active))
    assert "国文" in joined
    assert "701" in joined       # 教师视角格内显示班级


def test_room_export_ok_even_without_room(w):
    # 建一个没有课的场地也应导出成功(空课表)
    rid = w.client.post(f"/api/rooms{w.q}", json={"name": "101教室"}).json()["id"]
    r = w.client.get(f"/api/export/timetable{w.q}&view=room&target_id={rid}&fmt=xlsx")
    assert r.status_code == 200


# ── HTML(PDF 来源)─────────────────────────────────────────
def test_grid_html_has_chinese_and_layout(w):
    cid = w.classes["701"]
    grid, meta = tex.build_grid(w.db, w.sid, "class", cid)
    html = tex.grid_to_html(grid, meta)
    assert "国文" in html and "王师" in html
    assert "周三" in html and "第一节" in html
    assert "A4 portrait" in html
    assert meta.school_name in html


# ── 全校总表 / 批次 ─────────────────────────────────────────
def test_school_workbook_one_sheet_per_class(w):
    # 学期已有 701 与占位班 900 → 全校总表每班一分页
    r = w.client.get(f"/api/export/school.xlsx{w.q}")
    wb = _xlsx(r)
    assert len(wb.sheetnames) >= 2
    assert any("701" in s for s in wb.sheetnames)


def test_batch_zip_has_per_class_files(w):
    r = w.client.get(f"/api/export/batch.zip{w.q}")
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/zip"
    z = zipfile.ZipFile(io.BytesIO(r.content))
    names = z.namelist()
    assert any(n.endswith(".xlsx") for n in names)
    assert any("701" in n for n in names)


# ── 边界 / RBAC ─────────────────────────────────────────────
def test_export_unpublished_semester_404(env):
    client, db = env
    grant_grades(db, make_user(db, "s2", PW, roles=[Role.director]))
    client.post("/api/auth/login", json={"username": "s2", "password": PW})
    sid = client.post("/api/semesters", json={
        "academic_year": 2027, "term": 1, "template_key": "junior_high",
        "start_date": SEM_START.isoformat(), "end_date": SEM_END.isoformat(),
    }).json()["id"]
    cid = client.post(f"/api/class-units?semester_id={sid}",
                      json={"grade": 7, "name": "701", "track": "junior_high"}).json()["id"]
    r = client.get(f"/api/export/timetable?semester_id={sid}&view=class&target_id={cid}&fmt=xlsx")
    assert r.status_code == 404


def test_batch_60_classes_under_60s(db):
    """验收②:60 班批次导出 < 60 秒(Excel 产生为 CPU-bound,与资料库无关)。"""
    import time

    from app.models.timetable import Timetable, TimetableStatus
    from tests.fixtures import build_large_school

    fx = build_large_school(db, num_classes=60)
    db.add(Timetable(
        semester_id=fx.semester_id, name="正式课表",
        status=TimetableStatus.published.value))
    db.commit()

    t0 = time.perf_counter()
    data = tex.class_batch_zip(db, fx.semester_id)
    elapsed = time.perf_counter() - t0
    assert elapsed < 60, f"批次导出耗时 {elapsed:.1f}s"
    z = zipfile.ZipFile(io.BytesIO(data))
    assert len(z.namelist()) == 60


# ── 全校教师课表的总表与批次(用户回报 #22)──────────────
def test_school_workbook_teacher_view_has_a_sheet_per_teacher(w):
    """每位在职教师一个分页;没排到课的也要在(全校发放时不能漏人)。"""
    w.teacher("陈师", ["数学"])      # 有建教师但没排课
    w.teacher("林师", ["英文"])
    w.client.patch(f"/api/teachers/{w.teachers['林师']}", json={"name": "林师", "is_active": False})

    r = w.client.get(f"/api/export/school.xlsx{w.q}&view=teacher")
    wb = _xlsx(r)
    titles = "|".join(wb.sheetnames)
    assert "王师" in titles and "陈师" in titles      # 有课的、没课的都在
    assert "林师" not in titles                       # 离职的不列入
    assert len(wb.sheetnames) == 2

    # 内容:王师那张要有他周三第一节的国文
    ws = wb[next(n for n in wb.sheetnames if "王师" in n)]
    assert "国文" in " ".join(_cells(ws))


def test_batch_zip_teacher_view_has_one_file_per_teacher(w):
    w.teacher("陈师", ["数学"])
    r = w.client.get(f"/api/export/batch.zip{w.q}&view=teacher")
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    names = zipfile.ZipFile(io.BytesIO(r.content)).namelist()
    assert sorted(names) == ["王师.xlsx", "陈师.xlsx"]


def test_school_and_batch_default_to_class_view(w):
    """没带 view 时维持原本的班级版(旧的呼叫方式不受影响)。"""
    wb = _xlsx(w.client.get(f"/api/export/school.xlsx{w.q}"))
    assert any("701" in n for n in wb.sheetnames)
    names = zipfile.ZipFile(io.BytesIO(
        w.client.get(f"/api/export/batch.zip{w.q}").content)).namelist()
    assert all("年" in n for n in names)


def test_export_view_param_is_validated(w):
    assert w.client.get(f"/api/export/school.xlsx{w.q}&view=room").status_code == 422


def test_teacher_can_export_single_but_not_batch(w):
    make_user(w.db, "t", PW, roles=[Role.teacher])
    w.client.post("/api/auth/logout")
    w.client.post("/api/auth/login", json={"username": "t", "password": PW})
    cid = w.classes["701"]
    single = w.client.get(f"/api/export/timetable{w.q}&view=class&target_id={cid}&fmt=xlsx")
    assert single.status_code == 200          # 单一课表全校可查可导出
    assert w.client.get(f"/api/export/school.xlsx{w.q}").status_code == 403
    assert w.client.get(f"/api/export/batch.zip{w.q}").status_code == 403


# ── 打印/导出统一:值日行 ─────────────────────────────────────
def test_export_without_duty_has_no_duty_row(w):
    # 没排轮值的学期不画值日行,导出不凭空多一行。
    cid = w.classes["701"]
    r = w.client.get(f"/api/export/timetable{w.q}&view=class&target_id={cid}&fmt=xlsx")
    wb = _xlsx(r)
    first_cells = [row[0].value for row in wb.active.iter_rows()]
    assert "行政值日" not in first_cells
    grid, meta = tex.build_grid(w.db, w.sid, "class", cid)
    assert "行政值日" not in tex.grid_to_html(grid, meta)


def test_export_xlsx_and_html_have_duty_row(w):
    # 打印页已有值日行,导出也要有同一行:周三轮值王师,导出里周三格就是王师。
    from app.models.duty import DutyPattern
    w.db.add(DutyPattern(semester_id=w.sid, weekday=3, teacher_id=w.teachers["王师"]))
    w.db.commit()
    cid = w.classes["701"]
    r = w.client.get(f"/api/export/timetable{w.q}&view=class&target_id={cid}&fmt=xlsx")
    wb = _xlsx(r)
    duty_row = next(row for row in wb.active.iter_rows(values_only=True) if row[0] == "行政值日")
    assert duty_row[3] == "王师"  # 节次,周一,周二,周三
    grid, meta = tex.build_grid(w.db, w.sid, "class", cid)
    html = tex.grid_to_html(grid, meta)
    assert "行政值日" in html and "王师" in html
