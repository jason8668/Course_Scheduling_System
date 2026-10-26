"""课堂巡视情况记录表(v1.2.12):已发布课表叠上当天的调代课,印出那一天每班每节真正的样子。"""

import io
from datetime import timedelta

import pytest
from openpyxl import load_workbook

from app.models.user import Role
from tests.conftest import grant_grades, make_user
from tests.dates import SEM_END, SEM_START, WED
from tests.test_substitutions import _swap_first, _World

PW = "password123"


@pytest.fixture
def w(env):
    client, db = env
    grant_grades(db, make_user(db, "s", PW, roles=[Role.director]))
    client.post("/api/auth/login", json={"username": "s", "password": PW})
    sid = client.post("/api/semesters", json={
        "academic_year": 2026, "term": 1, "template_key": "junior_high",
        "start_date": SEM_START.isoformat(), "end_date": SEM_END.isoformat(),
    }).json()["id"]
    return _World(client, db, sid)


def _sheets(w, day=WED, to=None):
    q = f"{w.q}&date_from={day.isoformat()}" + (f"&date_to={to.isoformat()}" if to else "")
    return w.client.get(f"/api/patrol-sheets{q}")


def _cell(body, klass, ordinal, day=WED):
    """某班某节的格子(跨上午/下午与分页找)。"""
    for page in body["pages"]:
        if page["date"] != day.isoformat() or klass not in page["classes"]:
            continue
        for row in page["rows"]:
            if row["ordinal"] == ordinal:
                return row["cells"][page["classes"].index(klass)]
    raise AssertionError(f"找不到 {klass} 第 {ordinal} 节")


def test_sheet_follows_the_paper_form(w):
    """一天分上午(1~4 节)、下午(5~7 节)两张;早自习与午休不列;班级为栏。"""
    w.teacher("王师", ["国文"])
    w.place("王师", "国文", "701", 0)
    w.publish()

    r = _sheets(w)
    assert r.status_code == 200, r.json()
    body = r.json()
    assert "2026-2027学年第一学期" in body["title"]
    assert "授课情形" in body["legend"]
    assert [(p["first_ordinal"], p["last_ordinal"]) for p in body["pages"]] == [(1, 4), (5, 7)]
    morning = body["pages"][0]
    assert [row["name"] for row in morning["rows"]] == ["第一节", "第二节", "第三节", "第四节"]
    assert morning["table_name"] == ""          # 全校一套节次表:不必写是哪一套
    assert _cell(body, "701", 1) == {"subject": "国文", "teacher": "王师", "room": "", "note": ""}
    assert _cell(body, "701", 2) == {"subject": "", "teacher": "", "room": "", "note": ""}


def test_substitute_teacher_replaces_the_absent_one(w):
    """代课:那一天印代课老师并注记;其他日子(下周三)仍是原本的老师。"""
    w.teacher("王师", ["国文"])
    w.teacher("陈师", ["国文"])
    w.place("王师", "国文", "701", 0)
    w.publish()
    affected = w.leave("王师")
    assert _cell(_sheets(w).json(), "701", 1)["note"] == "请假待处理"

    status, _ = w.assign(affected[0]["id"], type="substitute",
                         handler_teacher_id=w.teachers["陈师"])
    assert status == 200
    cell = _cell(_sheets(w).json(), "701", 1)
    assert (cell["teacher"], cell["note"]) == ("陈师", "代课")

    next_week = WED + timedelta(days=7)
    cell = _cell(_sheets(w, next_week).json(), "701", 1, next_week)
    assert (cell["teacher"], cell["note"]) == ("王师", "")


def test_self_study_has_no_teacher(w):
    w.teacher("王师", ["国文"])
    w.place("王师", "国文", "701", 0)
    w.publish()
    affected = w.leave("王师")
    status, body = w.assign(affected[0]["id"], type="self_study")
    assert status == 200, body
    cell = _cell(_sheets(w).json(), "701", 1)
    assert (cell["teacher"], cell["note"]) == ("", "自习")


def test_swap_shows_on_both_days(w):
    """调课:请假那节印对调的老师,补课那节印回来补课的老师;两格都注记调课。"""
    w.teacher("王师", ["国文"])
    w.teacher("陈师", ["数学"])
    w.place("王师", "国文", "701", 0)                 # 周三第一节
    w.place("陈师", "数学", "701", 1, weekday=4)      # 周四第二节
    w.publish()
    affected = w.leave("王师")
    thu = WED + timedelta(days=1)
    _swap_first(w, affected[0]["id"], lambda o: o["date"] == thu.isoformat())

    body = _sheets(w, WED, thu).json()
    leave_cell = _cell(body, "701", 1, WED)
    makeup_cell = _cell(body, "701", 2, thu)
    # 同班互调:科目跟著老师走
    assert (leave_cell["subject"], leave_cell["teacher"], leave_cell["note"]) == (
        "数学", "陈师", "调课")
    assert (makeup_cell["subject"], makeup_cell["teacher"], makeup_cell["note"]) == (
        "国文", "王师", "调课")


def _group(w, name, classes):
    return w.client.post(f"/api/scheduling-units{w.q}", json={
        "name": name, "class_ids": [w.klass(c) for c in classes]}).json()["id"]


def _group_course(w, unit_id, subject, teacher, room=None):
    return w.client.post(f"/api/assignments{w.q}", json={
        "scheduling_unit_id": unit_id, "subject_id": w.subject(subject), "periods_per_week": 1,
        "teachers": [{"teacher_id": w.teachers[teacher]}], "block_rules": [], "room_id": room,
    }).json()["id"]


def _place_group(w, assignment_id, period_idx):
    """排入群组的任一门课,同群组的其他课会一起排进同一时段(H7)。"""
    r = w.client.post(f"/api/timetables/{w.tt}/entries", json={
        "course_assignment_id": assignment_id, "weekday": 3,
        "period_no": w.wed[period_idx]["period_no"], "span": 1})
    assert r.status_code == 201, r.json()


def test_small_group_lists_every_teacher_in_the_cell(w):
    """分组跑班(英语 A、B 两组):两班的格子都要看得到两位老师,不能只印其中一组。"""
    for t in ("英A师", "英B师"):
        w.teacher(t, ["英语"])
    unit = _group(w, "七年级英语分组", ["701", "702"])
    first = _group_course(w, unit, "英语", "英A师")
    _group_course(w, unit, "英语", "英B师")
    _place_group(w, first, 0)
    w.publish()

    body = _sheets(w).json()
    for klass in ("701", "702"):
        cell = _cell(body, klass, 1)
        assert (cell["subject"], cell["teacher"]) == ("英语", "英A师／英B师")
    assert body["group_lists"] == []


def test_big_group_gets_its_own_list(w):
    """社团:全校拆成很多组,格子塞不下,主表只写群组名称,另出一张分组巡视单。"""
    clubs = ["羽球社", "桌球社", "棒球社", "漫画社"]
    room_id = w.client.post(f"/api/rooms{w.q}", json={"name": "体育馆"}).json()["id"]
    unit = _group(w, "社团", ["701", "702"])
    ids = []
    for i, club in enumerate(clubs):
        w.teacher(f"师{i}", [club])
        ids.append(_group_course(w, unit, club, f"师{i}", room=room_id if i == 0 else None))
    _place_group(w, ids[0], 5)
    w.publish()

    body = _sheets(w).json()
    cell = _cell(body, "701", 6)
    assert (cell["subject"], cell["teacher"], cell["note"]) == ("社团", "共 4 组", "见分组巡视单")
    assert len(body["group_lists"]) == 1          # 两个班共用一张,不重复
    listing = body["group_lists"][0]
    assert (listing["group_name"], listing["period_name"]) == ("社团", "第六节")
    assert [(i["no"], i["subject"], i["teacher"]) for i in listing["items"]] == [
        (n, club, f"师{n - 1}") for n, club in enumerate(clubs, start=1)]
    assert listing["items"][0]["room"] == "体育馆"


def test_group_list_order_follows_the_assignments_not_the_entries(w):
    """分组单的编号照配课建立的顺序,与格位写入资料库的顺序无关。

    群组的格位是整批写入的,顺序取决于资料库怎么回传配课(PostgreSQL 不保证);
    这里刻意倒著写入,重现 CI 上「第一列变成棒球社」的情况。
    """
    from app.models.timetable import ScheduleEntry

    clubs = ["羽球社", "桌球社", "棒球社", "漫画社"]
    unit = _group(w, "社团", ["701", "702"])
    ids = []
    for i, club in enumerate(clubs):
        w.teacher(f"师{i}", [club])
        ids.append(_group_course(w, unit, club, f"师{i}"))
    for assignment_id in reversed(ids):
        w.db.add(ScheduleEntry(
            timetable_id=w.tt, course_assignment_id=assignment_id, weekday=3,
            period_no=w.wed[5]["period_no"], span=1))
    w.db.commit()
    w.publish()

    listing = _sheets(w).json()["group_lists"][0]
    assert [i["subject"] for i in listing["items"]] == clubs


def test_many_classes_are_split_across_pages(w):
    """班级超过一页的栏数时自动分页,每页上午/下午各一张。"""
    w.teacher("王师", ["国文"])
    w.place("王师", "国文", "701", 0)
    for n in range(13):
        w.klass(f"8{n:02d}")
    w.publish()

    pages = _sheets(w).json()["pages"]
    morning = [p for p in pages if p["first_ordinal"] == 1]
    assert [len(p["classes"]) for p in morning] == [12, 3]     # 701、900(占位班)+ 13 班
    assert len(pages) == 4


def test_days_without_lessons_are_skipped_and_errors_are_explained(w):
    w.teacher("王师", ["国文"])
    w.place("王师", "国文", "701", 0)

    assert _sheets(w).status_code == 404                       # 尚未发布
    w.publish()
    sat = WED + timedelta(days=3)
    r = _sheets(w, sat, sat + timedelta(days=1))               # 周末没有课
    assert r.status_code == 404 and "没有任何课" in r.json()["detail"]
    r = _sheets(w, WED, WED + timedelta(days=20))
    assert r.status_code == 404 and "最多" in r.json()["detail"]

    week = _sheets(w, WED - timedelta(days=2), WED + timedelta(days=4)).json()
    assert {p["date"] for p in week["pages"]} == {WED.isoformat()}   # 只有周三排了课


def test_legend_can_be_changed_by_admin(env):
    client, db = env
    make_user(db, "a", PW, roles=[Role.admin])
    client.post("/api/auth/login", json={"username": "a", "password": PW})
    assert "授课情形" in client.get("/api/settings/patrol").json()["legend"]
    r = client.put("/api/settings/patrol", json={"legend": "优、良、可"})
    assert r.json()["legend"] == "优、良、可"
    assert "授课情形" in client.put("/api/settings/patrol", json={"legend": ""}).json()["legend"]


# ── v1.2.14:导出 Excel ──────────────────────────────────────
def _xlsx(w, day=WED, to=None):
    q = f"{w.q}&date_from={day.isoformat()}" + (f"&date_to={to.isoformat()}" if to else "")
    r = w.client.get(f"/api/patrol-sheets.xlsx{q}")
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    return r, load_workbook(io.BytesIO(r.content))


def _texts(ws) -> list[str]:
    return [str(v) for row in ws.iter_rows(values_only=True) for v in row if v is not None]


def test_xlsx_has_a_sheet_per_half_day_with_the_same_content_as_the_print_page(w):
    """上午、下午各一个分页;代课老师与备注跟打印页一致,请假的老师不在表上。"""
    w.teacher("王师", ["国文"])
    w.teacher("陈师", ["国文"])
    w.place("王师", "国文", "701", 0)
    w.publish()
    affected = w.leave("王师")
    status, _ = w.assign(affected[0]["id"], type="substitute",
                         handler_teacher_id=w.teachers["陈师"])
    assert status == 200

    r, wb = _xlsx(w)
    day = f"{WED:%m-%d}"
    assert wb.sheetnames == [f"{day} 上午", f"{day} 下午"]
    assert f"UTF-8''%E8%AF%BE%E5%A0%82%E5%B7%A1%E8%A7%86%E6%83%85%E5%86%B5%E8%AE%B0%E5%BD%95%E8%A1%A8_{WED.isoformat()}.xlsx" in (
        r.headers["content-disposition"])

    ws = wb[f"{day} 上午"]
    texts = _texts(ws)
    assert "课堂巡视情况记录表" in texts[0] and "第 1~4 节" in texts[0]
    assert f"{WED.year}年{WED.month}月{WED.day}日" in texts[0]
    header = [c.value for c in ws[2]]
    assert header[0] == "节次" and "701" in header and "巡视" in header[-1]
    col = header.index("701") + 1
    # 第一节那一块:科目/教师、教室、授课情形、学生学习、备注
    assert [ws.cell(row=r_, column=2).value for r_ in range(3, 8)] == [
        "科目\n教师", "教室", "授课情形", "学生学习", "备注"]
    assert ws.cell(row=3, column=col).value == "国文\n陈师"
    assert ws.cell(row=7, column=col).value == "代课"
    assert ws.cell(row=5, column=col).value is None        # 授课情形留白给人填
    assert not any("王师" in t for t in texts)
    assert any("授课情形(填写代码)" in t for t in texts)     # 记录说明
    assert ws.page_setup.orientation == "landscape"         # 下载后不改也能直接印


def test_xlsx_puts_each_group_list_on_its_own_sheet(w):
    clubs = ["羽球社", "桌球社", "棒球社", "漫画社"]
    unit = _group(w, "社团", ["701", "702"])
    ids = []
    for i, club in enumerate(clubs):
        w.teacher(f"师{i}", [club])
        ids.append(_group_course(w, unit, club, f"师{i}"))
    _place_group(w, ids[0], 5)
    w.publish()

    _, wb = _xlsx(w)
    day = f"{WED:%m-%d}"
    assert wb.sheetnames == [f"{day} 上午", f"{day} 下午", f"{day} 社团 第六节"]
    ws = wb[f"{day} 社团 第六节"]
    assert [c.value for c in ws[3]] == ["编号", "名称", "教师", "地点", "上课状况", "学生表现"]
    assert [(ws.cell(row=4 + i, column=1).value, ws.cell(row=4 + i, column=2).value)
            for i in range(4)] == list(enumerate(clubs, start=1))
    # 主表那一格只写群组与组数
    assert any("社团\n共 4 组" == t for t in _texts(wb[f"{day} 下午"]))


def test_xlsx_week_export_and_errors(w):
    w.teacher("王师", ["国文"])
    w.place("王师", "国文", "701", 0)
    assert w.client.get(f"/api/patrol-sheets.xlsx{w.q}&date_from={WED}").status_code == 404
    w.publish()

    mon, sun = WED - timedelta(days=2), WED + timedelta(days=4)
    r, wb = _xlsx(w, mon, sun)
    assert len(wb.sheetnames) == 2                          # 只有周三排了课
    assert f"_{mon.isoformat()}_{sun:%m-%d}.xlsx" in r.headers["content-disposition"]


def test_patrol_sheets_filter_by_grade(w):
    """按年级打印:只印该年级的班;不存在的年级给 404。"""
    w.teacher("王师", ["国文"])
    w.teacher("陈师", ["国文"])
    w.place("王师", "国文", "701", 0)
    # 八年级班直接建班(grade=8),再让 _World 复用它排课(klass 会沿用已建的 id);
    # 用另一位老师,避免同一老师同节双班的硬约束冲突
    c801 = w.client.post(f"/api/class-units{w.q}", json={
        "grade": 8, "name": "801", "track": "junior_high"}).json()["id"]
    w.classes["801"] = c801
    w.place("陈师", "国文", "801", 0)
    w.publish()

    body = w.client.get(
        f"/api/patrol-sheets{w.q}&date_from={WED.isoformat()}&grade=7").json()
    classes7 = {c for p in body["pages"] for c in p["classes"]}
    assert "701" in classes7 and "801" not in classes7

    body = w.client.get(
        f"/api/patrol-sheets{w.q}&date_from={WED.isoformat()}&grade=8").json()
    classes8 = {c for p in body["pages"] for c in p["classes"]}
    assert classes8 == {"801"}

    r = w.client.get(f"/api/patrol-sheets{w.q}&date_from={WED.isoformat()}&grade=9")
    assert r.status_code == 404
