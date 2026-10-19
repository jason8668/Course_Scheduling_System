"""日薪代课与整批指派(v1.2.19,用户回报 #36)。

小学导师请整天假,常由一位代课老师整天接手、按日计酬:
- 计费方式多一个「日薪代课」;月结统计把这些节次从课时移到「日薪天数」。
- 整张假单(或其中一天)可以一次指派给同一位代课老师;派不成的节次跳过并说明原因。
"""

import io

import pytest
from openpyxl import load_workbook

from app.models.user import Role
from tests.conftest import make_user
from tests.dates import SEM_END, SEM_START, THU, WED
from tests.test_substitutions import _World

PW = "password123"
DAILY = "日薪代课"


@pytest.fixture
def w(env):
    client, db = env
    make_user(db, "s", PW, roles=[Role.scheduler])
    client.post("/api/auth/login", json={"username": "s", "password": PW})
    sid = client.post("/api/semesters", json={
        "academic_year": 115, "term": 1, "template_key": "junior_high",
        "start_date": SEM_START.isoformat(), "end_date": SEM_END.isoformat(),
    }).json()["id"]
    return _World(client, db, sid)


def _stats(w, day=WED, **params):
    qs = "".join(f"&{k}={v}" for k, v in params.items())
    return w.client.get(
        f"/api/substitution-stats{w.q}&year={day.year}&month={day.month}{qs}").json()


def _summary(w, name, day=WED):
    return next(s for s in _stats(w, day)["summaries"] if s["teacher_name"] == name)


def _leave(w, teacher, start, end):
    """跨日假单;回传 (leave_id, 受影响节次)。"""
    r = w.client.post(f"/api/leaves{w.q}", json={
        "teacher_id": w.teachers[teacher], "leave_type": "sick",
        "start_date": start.isoformat(), "end_date": end.isoformat()})
    assert r.status_code == 201, r.json()
    return r.json()["id"], r.json()["affected_periods"]


def _batch(w, leave_id, teacher, **body):
    return w.client.post(f"/api/leaves/{leave_id}/substitutions/batch",
                         json={"handler_teacher_id": w.teachers[teacher], **body})


def _homeroom_day(w, periods=3, weekday=3):
    """王师(导师)某一天连上几节 701 的课(包班:每节不同科目,不会撞每日上限)。"""
    for i in range(periods):
        w.place("王师", ("国文", "数学", "生活")[i], "701", i, weekday=weekday)


# ── 计费方式 ──────────────────────────
def test_daily_is_offered_as_a_funding_option(w):
    options = w.client.get("/api/substitution-funding-sources").json()
    assert DAILY in options


# ── 统计:课时与日薪分开 ────────────────
def test_daily_periods_count_as_days_not_hours(w):
    """陈师日薪代王师一整天 3 节 → 日薪 1 天,课时计费 0 节;代课节数仍是 3。"""
    w.teacher("王师", ["国文"])
    w.teacher("陈师", ["国文"])
    _homeroom_day(w)
    w.publish()
    for ap in w.leave("王师"):
        code, _ = w.assign(ap["id"], type="substitute", handler_teacher_id=w.teachers["陈师"],
                           funding_source=DAILY)
        assert code == 200

    chen = _summary(w, "陈师")
    assert chen["handled_count"] == 3
    assert chen["billable_count"] == 0
    assert chen["daily_days"] == 1 and chen["daily_periods"] == 3
    assert chen["daily_shared_days"] == 0
    details = _stats(w)["details"]
    assert {d["pay_kind"] for d in details} == {"daily"}
    assert not any(d["daily_shared"] for d in details)


def test_hourly_and_daily_are_listed_separately_for_the_same_teacher(w):
    """同一位代课老师这个月既有课时也有日薪:两个数字互不相混。"""
    w.teacher("王师", ["国文"])
    w.teacher("李师", ["国文"])
    w.teacher("陈师", ["国文"])
    _homeroom_day(w, periods=2)
    w.place("李师", "国文", "702", 3)
    w.publish()
    for ap in w.leave("王师"):
        w.assign(ap["id"], type="substitute", handler_teacher_id=w.teachers["陈师"],
                 funding_source=DAILY)
    li = w.leave("李师")
    w.assign(li[0]["id"], type="substitute", handler_teacher_id=w.teachers["陈师"])  # 公费课时

    chen = _summary(w, "陈师")
    assert chen["handled_count"] == 3
    assert chen["billable_count"] == 1
    assert chen["daily_days"] == 1 and chen["daily_periods"] == 2
    kinds = sorted(d["pay_kind"] for d in _stats(w)["details"])
    assert kinds == ["daily", "daily", "hourly"]


def test_schools_without_daily_substitutes_see_the_same_numbers_as_before(w):
    w.teacher("王师", ["国文"])
    w.teacher("陈师", ["国文"])
    _homeroom_day(w, periods=2)
    w.publish()
    ap = w.leave("王师")
    w.assign(ap[0]["id"], type="substitute", handler_teacher_id=w.teachers["陈师"])
    w.assign(ap[1]["id"], type="merge", handler_teacher_id=w.teachers["陈师"])

    chen = _summary(w, "陈师")
    assert (chen["handled_count"], chen["billable_count"]) == (2, 1)
    assert (chen["daily_days"], chen["daily_periods"], chen["daily_shared_days"]) == (0, 0, 0)
    assert sorted(d["pay_kind"] for d in _stats(w)["details"]) == ["hourly", "none"]


def test_day_shared_by_two_daily_substitutes_is_flagged_for_both(w):
    """找不到同一位代课老师、一天分给两位:各算 1 天并标示,怎么给付由学校决定。"""
    w.teacher("王师", ["国文"])
    w.teacher("陈师", ["国文"])
    w.teacher("林师", ["国文"])
    _homeroom_day(w)
    w.publish()
    ap = w.leave("王师")
    w.assign(ap[0]["id"], type="substitute", handler_teacher_id=w.teachers["陈师"],
             funding_source=DAILY)
    w.assign(ap[1]["id"], type="substitute", handler_teacher_id=w.teachers["陈师"],
             funding_source=DAILY)
    w.assign(ap[2]["id"], type="substitute", handler_teacher_id=w.teachers["林师"],
             funding_source=DAILY)

    chen, lin = _summary(w, "陈师"), _summary(w, "林师")
    assert (chen["daily_days"], chen["daily_periods"], chen["daily_shared_days"]) == (1, 2, 1)
    assert (lin["daily_days"], lin["daily_periods"], lin["daily_shared_days"]) == (1, 1, 1)
    assert all(d["daily_shared"] for d in _stats(w)["details"])
    # 只查其中一位(教师个人查询也是这样筛)仍然看得到「这天是分担的」
    only = _stats(w, teacher_id=w.teachers["林师"])
    assert only["summaries"][0]["daily_shared_days"] == 1


def test_daily_plus_hourly_on_the_same_day_is_not_a_shared_day(w):
    """另一位是课时代课就不算「日薪分担」——日薪那位还是完整算一天。"""
    w.teacher("王师", ["国文"])
    w.teacher("陈师", ["国文"])
    w.teacher("林师", ["国文"])
    _homeroom_day(w, periods=2)
    w.publish()
    ap = w.leave("王师")
    w.assign(ap[0]["id"], type="substitute", handler_teacher_id=w.teachers["陈师"],
             funding_source=DAILY)
    w.assign(ap[1]["id"], type="substitute", handler_teacher_id=w.teachers["林师"])

    assert _summary(w, "陈师")["daily_shared_days"] == 0
    assert _summary(w, "林师")["billable_count"] == 1


def test_export_has_daily_columns(w):
    w.teacher("王师", ["国文"])
    w.teacher("陈师", ["国文"])
    w.teacher("林师", ["国文"])
    _homeroom_day(w, periods=2)
    w.publish()
    ap = w.leave("王师")
    w.assign(ap[0]["id"], type="substitute", handler_teacher_id=w.teachers["陈师"],
             funding_source=DAILY)
    w.assign(ap[1]["id"], type="substitute", handler_teacher_id=w.teachers["林师"],
             funding_source=DAILY)

    r = w.client.get(f"/api/substitution-stats/export{w.q}&year={WED.year}&month={WED.month}")
    wb = load_workbook(io.BytesIO(r.content))
    summary = [[c.value for c in row] for row in wb["汇总"].iter_rows()]
    assert summary[0] == ["教师", "代课节数", "课时计费节数", "日薪天数", "日薪节数", "备注"]
    chen = next(row for row in summary if row[0] == "陈师")
    assert chen[1:5] == [1, 0, 1, 1] and "分担" in chen[5]
    detail = [[c.value for c in row] for row in wb["明细"].iter_rows()]
    assert detail[0][-2:] == ["计酬", "备注"]
    assert detail[1][-2] == "日薪" and "分担" in detail[1][-1]


# ── 整批指派 ──────────────────────────
def test_batch_assigns_every_pending_period_of_the_leave(w):
    w.teacher("王师", ["国文"])
    w.teacher("陈师", ["国文"])
    _homeroom_day(w, periods=3, weekday=3)
    _homeroom_day(w, periods=2, weekday=4)
    w.publish()
    leave_id, ap = _leave(w, "王师", WED, THU)
    assert len(ap) == 5

    r = _batch(w, leave_id, "陈师", funding_source=DAILY)
    assert r.status_code == 200, r.text
    assert r.json() == {"assigned": 5, "skipped": []}

    leave = next(x for x in w.client.get(f"/api/leaves{w.q}").json() if x["id"] == leave_id)
    assert {p["status"] for p in leave["affected_periods"]} == {"resolved"}
    assert {p["handler_name"] for p in leave["affected_periods"]} == {"陈师"}
    # 两天 → 日薪 2 天(两天可能跨月,各月分开查再加总)
    months = {(WED.year, WED.month): WED, (THU.year, THU.month): THU}
    assert sum(_summary(w, "陈师", d)["daily_days"] for d in months.values()) == 2


def test_batch_can_be_limited_to_one_day(w):
    w.teacher("王师", ["国文"])
    w.teacher("陈师", ["国文"])
    _homeroom_day(w, periods=3, weekday=3)
    _homeroom_day(w, periods=2, weekday=4)
    w.publish()
    leave_id, _ = _leave(w, "王师", WED, THU)

    r = _batch(w, leave_id, "陈师", date=THU.isoformat())
    assert r.json()["assigned"] == 2
    leave = next(x for x in w.client.get(f"/api/leaves{w.q}").json() if x["id"] == leave_id)
    by_day = {p["date"]: p["status"] for p in leave["affected_periods"]}
    assert by_day == {WED.isoformat(): "pending", THU.isoformat(): "resolved"}


def test_batch_skips_periods_the_substitute_cannot_take_and_says_why(w):
    """代课老师第二节自己有课:那一节跳过并说明,其余照派——组长再另外找人。"""
    w.teacher("王师", ["国文"])
    w.teacher("陈师", ["国文"])
    _homeroom_day(w, periods=3)
    w.place("陈师", "国文", "702", 1)  # 陈师周三第二节有自己的课
    w.publish()
    leave_id, ap = _leave(w, "王师", WED, WED)

    r = _batch(w, leave_id, "陈师")
    body = r.json()
    assert body["assigned"] == 2
    assert [s["affected_period_id"] for s in body["skipped"]] == [ap[1]["id"]]
    skip = body["skipped"][0]
    assert skip["period_name"] == "第二节" and skip["class_names"] == "701"
    assert "陈师" in skip["reason"]
    leave = next(x for x in w.client.get(f"/api/leaves{w.q}").json() if x["id"] == leave_id)
    assert [p["status"] for p in leave["affected_periods"]] == ["resolved", "pending", "resolved"]


def test_batch_leaves_already_handled_periods_alone(w):
    w.teacher("王师", ["国文"])
    w.teacher("陈师", ["国文"])
    w.teacher("林师", ["国文"])
    _homeroom_day(w, periods=3)
    w.publish()
    leave_id, ap = _leave(w, "王师", WED, WED)
    w.assign(ap[0]["id"], type="substitute", handler_teacher_id=w.teachers["林师"])

    assert _batch(w, leave_id, "陈师").json()["assigned"] == 2
    leave = next(x for x in w.client.get(f"/api/leaves{w.q}").json() if x["id"] == leave_id)
    assert [p["handler_name"] for p in leave["affected_periods"]] == ["林师", "陈师", "陈师"]


def test_batch_with_nothing_left_to_assign_is_rejected(w):
    w.teacher("王师", ["国文"])
    w.teacher("陈师", ["国文"])
    _homeroom_day(w, periods=1)
    w.publish()
    leave_id, _ = _leave(w, "王师", WED, WED)
    assert _batch(w, leave_id, "陈师").status_code == 200

    again = _batch(w, leave_id, "陈师")
    assert again.status_code == 409
    assert "没有待处理的节次" in again.json()["detail"]
    assert _batch(w, 999_999, "陈师").status_code == 404


def test_batch_cannot_assign_the_absent_teacher_to_their_own_leave(w):
    """每一节都因同一个原因派不成 → 全部列在 skipped,没有任何节次被改动。"""
    w.teacher("王师", ["国文"])
    _homeroom_day(w, periods=2)
    w.publish()
    leave_id, _ = _leave(w, "王师", WED, WED)

    body = _batch(w, leave_id, "王师").json()
    assert body["assigned"] == 0 and len(body["skipped"]) == 2
    assert all("请假教师" in s["reason"] for s in body["skipped"])
