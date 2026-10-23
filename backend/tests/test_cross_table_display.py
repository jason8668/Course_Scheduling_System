"""跨节次表教师的「显示」:初中小(初中 45 分、小学 40 分)同一位老师两部都有课。

冲堂判定本来就以墙钟时间为准(D7,见 test_timetables);这里验的是用户看到的东西——
教师课表与通知单的格线只能照一套节次表画,另一套的课必须标出实际时间,
冲堂讯息也要写出对方那一节真正的时间。
"""

import io

import pytest
from openpyxl import load_workbook

from app.models.user import Role
from tests.conftest import make_user
from tests.dates import SEM_END, SEM_START, WED
from tests.test_substitutions import _World

PW = "password123"


@pytest.fixture
def k9(env):
    """初中小:默认初中节次表 + 小学节次表;林师在 701(初中)与 601(小学)都有英语。"""
    client, db = env
    make_user(db, "s", PW, roles=[Role.scheduler])
    client.post("/api/auth/login", json={"username": "s", "password": PW})
    sid = client.post("/api/semesters", json={
        "academic_year": 2026, "term": 1, "template_key": "junior_high",
        "start_date": SEM_START.isoformat(), "end_date": SEM_END.isoformat(),
    }).json()["id"]
    w = _World(client, db, sid)
    table = client.post(f"/api/semesters/{sid}/period-tables",
                        json={"name": "小学节次表", "template_key": "elementary"}).json()
    w.classes["601"] = client.post(f"/api/class-units{w.q}", json={
        "grade": 6, "name": "601", "track": "elementary", "period_table_id": table["id"],
    }).json()["id"]
    w.teacher("林师", ["英语"])
    w.teacher("陈师", ["英语"])
    w.place("林师", "英语", "701", 0)   # 初中第一节 08:20–09:05
    w.place("林师", "英语", "601", 3)   # 小学第四节 11:10–11:50(初中第四节是 11:15–12:00)
    return w


def test_conflict_message_gives_the_other_lessons_real_time(k9):
    """小学第一节 08:30 撞到初中第一节 08:20–09:05:讯息要写出那堂课实际的时间。"""
    w = k9
    slots = [p for p in w.client.get(
        f"/api/class-units/{w.classes['601']}/period-table").json()["periods"]
        if p["weekday"] == 3 and p["type"] == "regular"]
    a = w.client.post(f"/api/assignments{w.q}", json={
        "class_id": w.classes["601"], "subject_id": w.subject("英语"), "periods_per_week": 1,
        "teachers": [{"teacher_id": w.teachers["林师"]}], "block_rules": [],
    }).json()
    r = w.client.post(f"/api/timetables/{w.tt}/entries", json={
        "course_assignment_id": a["id"], "weekday": 3,
        "period_no": slots[0]["period_no"], "span": 1})
    assert r.status_code == 409
    messages = [c["message"] for c in r.json()["detail"]["conflicts"]]
    assert any("701 班英语(08:20–09:05)" in m for m in messages), messages


def test_teacher_timetable_marks_the_real_time_of_other_table_lessons(k9):
    w = k9
    w.publish()
    r = w.client.get(
        f"/api/export/timetable{w.q}&view=teacher&target_id={w.teachers['林师']}&fmt=xlsx")
    assert r.status_code == 200, r.text
    cells = [str(v) for row in load_workbook(io.BytesIO(r.content)).active.iter_rows(
        values_only=True) for v in row if v]
    c601 = next(c for c in cells if "601" in c)
    c701 = next(c for c in cells if "701" in c)
    assert "11:10–11:50" in c601          # 小学的课:标出实际时间
    assert "–" not in c701                # 初中的课就在自己的格线上,不必标


def test_class_timetable_needs_no_note(k9):
    """班级课表照该班自己的节次表画,时间本来就对。"""
    w = k9
    w.publish()
    r = w.client.get(
        f"/api/export/timetable{w.q}&view=class&target_id={w.classes['601']}&fmt=xlsx")
    cells = [str(v) for row in load_workbook(io.BytesIO(r.content)).active.iter_rows(
        values_only=True) for v in row if v]
    assert "–" not in next(c for c in cells if "林师" in c)


def test_substitute_slip_marks_the_real_time_of_other_table_lessons(k9):
    """陈师同一张代课单上有初中与小学的课:格线照第一节(初中)画,小学那格要标时间。"""
    w = k9
    w.publish()
    affected = w.leave("林师")
    for ap in affected:
        status, _ = w.assign(ap["id"], type="substitute", handler_teacher_id=w.teachers["陈师"])
        assert status == 200
    q = "&".join(f"affected_period_ids={ap['id']}" for ap in affected)
    body = w.client.get(f"/api/substitute-slips{w.q}&{q}").json()

    teacher = next(s for s in body["slips"] if s["kind"] == "teacher")
    notes = {c["actor"]: c["time_note"] for c in teacher["weeks"][0]["cells"]}
    assert notes == {"701[代]": "", "601[代]": "11:10–11:50"}
    assert WED.isoformat() == teacher["weeks"][0]["cells"][0]["date"]

    for slip in (s for s in body["slips"] if s["kind"] == "class"):
        assert all(c["time_note"] == "" for c in slip["weeks"][0]["cells"])   # 班级单照自己的格线


# ── 完全中学:两部节数不同(初中七节、高中八节)──────────────────
@pytest.fixture
def k12(env):
    """默认初中节次表(七节)+ 高中部节次表(八节);林师在 901 与高一甲都有英语。"""
    client, db = env
    make_user(db, "s", PW, roles=[Role.scheduler])
    client.post("/api/auth/login", json={"username": "s", "password": PW})
    sid = client.post("/api/semesters", json={
        "academic_year": 2026, "term": 1, "template_key": "junior_high",
        "start_date": SEM_START.isoformat(), "end_date": SEM_END.isoformat(),
    }).json()["id"]
    w = _World(client, db, sid)
    table = client.post(f"/api/semesters/{sid}/period-tables",
                        json={"name": "高中部节次表", "template_key": "senior_high"}).json()
    w.classes["高一甲"] = client.post(f"/api/class-units{w.q}", json={
        "grade": 10, "name": "高一甲", "track": "senior_high", "period_table_id": table["id"],
    }).json()["id"]
    w.teacher("林师", ["英语"])
    w.teacher("陈师", ["英语"])
    w.place("林师", "英语", "901", 0)      # 初中第一节
    w.place("林师", "英语", "高一甲", 7)   # 高中第八节 16:20–17:10:初中节次表没有这一列
    w.publish()
    return w


def test_teacher_timetable_adds_the_row_the_default_table_lacks(k12):
    """高中第八节的课不能因为初中格线只有七列就从教师课表上消失。"""
    w = k12
    r = w.client.get(
        f"/api/export/timetable{w.q}&view=teacher&target_id={w.teachers['林师']}&fmt=xlsx")
    assert r.status_code == 200, r.text
    cells = [str(v) for row in load_workbook(io.BytesIO(r.content)).active.iter_rows(
        values_only=True) for v in row if v]
    assert "第八节" in cells
    eighth = next(c for c in cells if "高一甲" in c)
    assert "16:20–17:10" in eighth


def test_teacher_without_cross_table_lessons_gets_no_extra_row(k12):
    """只在初中上课的老师,课表维持七列,不会多一列空的第八节。"""
    w = k12
    w2_teacher = w.teachers["陈师"]
    r = w.client.get(f"/api/export/timetable{w.q}&view=teacher&target_id={w2_teacher}&fmt=xlsx")
    cells = [str(v) for row in load_workbook(io.BytesIO(r.content)).active.iter_rows(
        values_only=True) for v in row if v]
    assert "第七节" in cells and "第八节" not in cells


def test_substitute_slip_uses_the_table_with_more_rows(k12):
    """教师代课单的格线照节数多的那套(高中八节)画,第八节那格才有地方放。"""
    w = k12
    affected = w.leave("林师")
    for ap in affected:
        status, _ = w.assign(ap["id"], type="substitute", handler_teacher_id=w.teachers["陈师"])
        assert status == 200
    q = "&".join(f"affected_period_ids={ap['id']}" for ap in affected)
    body = w.client.get(f"/api/substitute-slips{w.q}&{q}").json()

    teacher = next(s for s in body["slips"] if s["kind"] == "teacher")
    assert len(teacher["rows"]) == 8
    cells = {c["actor"]: (c["ordinal"], c["time_note"]) for c in teacher["weeks"][0]["cells"]}
    # 格线是高中的:高中那格不必标时间,初中那格标实际时间
    assert cells == {"901[代]": (1, "08:20–09:05"), "高一甲[代]": (8, "")}


# ── 第八节:有的学校有、有的没有,初中/高中各种组合 ──────────────
def _add_eighth_period(w, table, start, end):
    """在节次表最后加一列第八节(等同「编辑节次表 → 新增节次列」)。时间可以不填。"""
    periods = [{k: p[k] for k in ("weekday", "period_no", "name", "start_time", "end_time", "type")}
               for p in table["periods"]]
    for wd in range(1, 6):
        periods.append({"weekday": wd, "period_no": 10, "name": "第八节",
                        "start_time": start, "end_time": end, "type": "regular"})
    r = w.client.put(f"/api/period-tables/{table['id']}/periods", json=periods)
    assert r.status_code == 200, r.text


def _k12_with_junior_eighth(env, start, end):
    client, db = env
    make_user(db, "s", PW, roles=[Role.scheduler])
    client.post("/api/auth/login", json={"username": "s", "password": PW})
    sem = client.post("/api/semesters", json={
        "academic_year": 2026, "term": 1, "template_key": "junior_high",
        "start_date": SEM_START.isoformat(), "end_date": SEM_END.isoformat(),
    }).json()
    w = _World(client, db, sem["id"])
    _add_eighth_period(w, sem["period_tables"][0], start, end)
    table = client.post(f"/api/semesters/{sem['id']}/period-tables",
                        json={"name": "高中部节次表", "template_key": "senior_high"}).json()
    w.classes["高一甲"] = client.post(f"/api/class-units{w.q}", json={
        "grade": 10, "name": "高一甲", "track": "senior_high", "period_table_id": table["id"],
    }).json()["id"]
    w.teacher("林师", ["英语"])
    return w


def _try_place(w, klass, period_idx, weekday=3):
    slots = [p for p in w.client.get(
        f"/api/class-units/{w.klass(klass)}/period-table").json()["periods"]
        if p["weekday"] == weekday and p["type"] == "regular"]
    a = w.client.post(f"/api/assignments{w.q}", json={
        "class_id": w.klass(klass), "subject_id": w.subject("英语"), "periods_per_week": 1,
        "teachers": [{"teacher_id": w.teachers["林师"]}], "block_rules": [],
    }).json()
    return w.client.post(f"/api/timetables/{w.tt}/entries", json={
        "course_assignment_id": a["id"], "weekday": weekday,
        "period_no": slots[period_idx]["period_no"], "span": 1})


def test_both_divisions_have_an_eighth_period(env):
    """初中、高中都有第八节(16:05–16:50 与 16:20–17:10):同一天会挡,课表只有一列第八节。"""
    w = _k12_with_junior_eighth(env, "16:05:00", "16:50:00")
    assert _try_place(w, "901", 7).status_code == 201
    r = _try_place(w, "高一甲", 7)
    assert r.status_code == 409
    assert "901 班英语(16:05–16:50)" in r.json()["detail"]["conflicts"][0]["message"]
    assert _try_place(w, "高一甲", 7, weekday=4).status_code == 201   # 换一天就可以

    w.publish()
    r = w.client.get(
        f"/api/export/timetable{w.q}&view=teacher&target_id={w.teachers['林师']}&fmt=xlsx")
    cells = [str(v) for row in load_workbook(io.BytesIO(r.content)).active.iter_rows(
        values_only=True) for v in row if v]
    assert cells.count("第八节") == 1
    assert "16:20–17:10" in next(c for c in cells if "高一甲" in c)
    assert "–" not in next(c for c in cells if "901" in c)


def test_period_without_times_still_blocks_the_same_period_number(env):
    """学校自己加的第八节没填起讫时间:无从比时间,退回节次号——不能让老师同时段两堂课。"""
    w = _k12_with_junior_eighth(env, None, None)
    assert _try_place(w, "901", 7).status_code == 201
    r = _try_place(w, "高一甲", 7)
    assert r.status_code == 409, r.json()
    assert _try_place(w, "高一甲", 6).status_code == 201   # 不同节次号照常可排
