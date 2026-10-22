"""Excel 导入:范本产生、逐列验证、交易式入库。

范本固定三列表头:第 1 列栏名、第 2 列说明、第 3 列范例;
导入时自第 4 列起读取资料(前三列自动略过)。
验证采「全对才写入」:任一列有误即回报所有错误,资料库零写入。
"""

import io
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.security import hash_password
from app.core.validators import is_valid_email
from app.models.assignment import AssignmentTeacher, BlockRule, CourseAssignment
from app.models.basedata import (
    ClassTrack,
    ClassUnit,
    RoomType,
    Subject,
    Teacher,
)
from app.models.period import PeriodTable
from app.models.user import Role, User, UserRole
from app.services.assignments import (
    DomainError,
    assert_within_overtime_limit,
    get_or_create_single_unit,
)

HEADER_ROWS = 3  # 栏名 + 说明 + 范例

ROOM_TYPE_BY_LABEL = {
    "普通教室": RoomType.normal,
    "专科教室": RoomType.special,
    "实习工场": RoomType.workshop,
    "户外": RoomType.outdoor,
}
TRACK_BY_LABEL = {
    "小学": ClassTrack.elementary,
    "初中": ClassTrack.junior_high,
    "普通型高中": ClassTrack.senior_high,
    "综合型高中": ClassTrack.comprehensive,
    "技术型高中": ClassTrack.vocational,
}

# 每个实体的范本栏位:(栏名, 说明, 范例)
TEMPLATE_DEFS: dict[str, dict] = {
    "subjects": {
        "sheet": "科目",
        "columns": [
            ("名称", "必填", "数学"),
            ("领域/群别", "选填", "数学领域"),
            ("需要场地类型", "选填:普通教室/专科教室/实习工场/户外", "普通教室"),
            ("默认连堂", "选填,数字 1-8,默认 1", "1"),
        ],
    },
    "teachers": {
        "sheet": "教师",
        "columns": [
            ("姓名", "必填", "王小明"),
            ("身分末四码", "选填,4 码,用于辨识同名教师;更新既有教师时以姓名＋末四码比对", "1234"),
            ("任教科目", "选填,多科以、分隔;需为已建立的科目", "数学、物理"),
            ("基本课时", "选填,数字", "20"),
            ("行政职称", "选填;更新时填「无」代表卸任", "教学组长"),
            ("行政减课", "选填,数字", "4"),
            ("外聘", "选填:是/否,默认否", "否"),
            ("登录帐号", "选填,勾选建立帐号时使用", "wang001"),
            ("Email", "选填,调代课通知寄送用", "wang@example.edu.tw"),
            ("手机", "选填,人工联络用", "0912345678"),
        ],
    },
    "classes": {
        "sheet": "班级",
        "columns": [
            ("年级", "必填,数字 1-12", "1"),
            ("班名", "必填", "甲"),
            ("学制", "必填:小学/初中/普通型高中/综合型高中/技术型高中", "技术型高中"),
            ("群科", "选填(技高填写)", "机械科"),
            ("班主任", "选填,需为已建立的教师姓名", "王小明"),
            ("人数", "选填,数字", "35"),
            ("节次表", "选填,需为已建立的节次表名称;空白则用学期默认", "高中部节次表"),
        ],
    },
    "assignments": {
        "sheet": "配课",
        "columns": [
            ("班级", "必填,需为已建立的班名(单班配课;跑班群组请于界面建立)", "甲"),
            ("科目", "必填,需为已建立的科目", "国文"),
            ("教师", "必填,多位以、分隔,第一位为主教", "王小明、李协同"),
            ("每周节数", "必填,数字", "5"),
            ("连堂长度", "选填,2-4;与连堂次数成对填写", "2"),
            ("连堂次数", "选填,数字", "1"),
            ("场地类型", "选填:普通教室/专科教室/实习工场/户外", "专科教室"),
        ],
    },
}


@dataclass
class ImportResult:
    imported: int = 0
    updated: int = 0          # 更新既有资料的笔数(目前只有教师导入会用到)
    errors: list[str] = field(default_factory=list)


def build_template(entity: str) -> bytes:
    cfg = TEMPLATE_DEFS[entity]
    wb = Workbook()
    ws = wb.active
    ws.title = cfg["sheet"]
    ws.append([c[0] for c in cfg["columns"]])
    ws.append([c[1] for c in cfg["columns"]])
    ws.append([c[2] for c in cfg["columns"]])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(row: tuple, i: int) -> str | None:
    if i >= len(row) or row[i] is None:
        return None
    text = str(row[i]).strip()
    return text or None


# 更新模式下,这些值代表「清空这一栏」(例如卸任行政职)
_CLEAR_TOKENS = {"无", "-", "—", "none", "None"}


def _parse_int(value: str | None) -> int | None:
    if value is None:
        return None
    try:
        return int(float(value))  # 容忍 Excel 把数字读成 20.0
    except ValueError as err:
        raise ValueError(f"「{value}」不是有效数字") from err


def _data_rows(file_bytes: bytes):
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active
    for idx, row in enumerate(
        ws.iter_rows(min_row=HEADER_ROWS + 1, values_only=True), start=HEADER_ROWS + 1
    ):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        yield idx, row


# ── 各实体导入 ────────────────────────
def _import_subjects(db: Session, semester_id: int, file_bytes: bytes) -> ImportResult:
    result = ImportResult()
    pending: list[Subject] = []
    for idx, row in _data_rows(file_bytes):
        name = _cell(row, 0)
        if not name:
            result.errors.append(f"第 {idx} 列:名称必填")
            continue
        room_label = _cell(row, 2)
        room_type = None
        if room_label:
            if room_label not in ROOM_TYPE_BY_LABEL:
                result.errors.append(f"第 {idx} 列:场地类型「{room_label}」无效")
                continue
            room_type = ROOM_TYPE_BY_LABEL[room_label].value
        try:
            block = _parse_int(_cell(row, 3)) or 1
        except ValueError as e:
            result.errors.append(f"第 {idx} 列:默认连堂 {e}")
            continue
        pending.append(
            Subject(
                semester_id=semester_id, name=name, domain=_cell(row, 1),
                required_room_type=room_type, default_block_size=block,
            )
        )
    if result.errors:
        return result
    db.add_all(pending)
    db.commit()
    result.imported = len(pending)
    return result


def _import_classes(db: Session, semester_id: int, file_bytes: bytes) -> ImportResult:
    result = ImportResult()
    teachers = {
        t.name: t.id
        for t in db.scalars(select(Teacher).where(Teacher.semester_id == semester_id))
    }
    period_tables = {
        pt.name: pt.id
        for pt in db.scalars(
            select(PeriodTable).where(PeriodTable.semester_id == semester_id)
        )
    }
    # 同学期班名唯一(M6-5)。档案内重复、或与既有班级重复,都要当场说是哪一列,
    # 不能让它撞上 DB 约束变成一句看不懂的错误。
    existing_names = set(
        db.scalars(
            select(ClassUnit.name).where(ClassUnit.semester_id == semester_id)
        )
    )
    pending: list[ClassUnit] = []
    for idx, row in _data_rows(file_bytes):
        try:
            grade = _parse_int(_cell(row, 0))
        except ValueError as e:
            result.errors.append(f"第 {idx} 列:年级 {e}")
            continue
        name = _cell(row, 1)
        track_label = _cell(row, 2)
        if not grade or not name or not track_label:
            result.errors.append(f"第 {idx} 列:年级、班名、学制皆必填")
            continue
        if name in existing_names:
            result.errors.append(f"第 {idx} 列:班名「{name}」在本学期重复")
            continue
        existing_names.add(name)
        if track_label not in TRACK_BY_LABEL:
            result.errors.append(f"第 {idx} 列:学制「{track_label}」无效")
            continue
        homeroom_name = _cell(row, 4)
        homeroom_id = None
        if homeroom_name:
            if homeroom_name not in teachers:
                result.errors.append(f"第 {idx} 列:班主任「{homeroom_name}」不存在")
                continue
            homeroom_id = teachers[homeroom_name]
        table_name = _cell(row, 6)
        table_id = None
        if table_name:
            if table_name not in period_tables:
                result.errors.append(f"第 {idx} 列:节次表「{table_name}」不存在")
                continue
            table_id = period_tables[table_name]
        try:
            count = _parse_int(_cell(row, 5))
        except ValueError as e:
            result.errors.append(f"第 {idx} 列:人数 {e}")
            continue
        pending.append(
            ClassUnit(
                semester_id=semester_id, grade=grade, name=name,
                track=TRACK_BY_LABEL[track_label].value, department=_cell(row, 3),
                student_count=count, homeroom_teacher_id=homeroom_id,
                period_table_id=table_id,
            )
        )
    if result.errors:
        return result
    db.add_all(pending)
    db.commit()
    result.imported = len(pending)
    return result


def _import_teachers(
    db: Session,
    semester_id: int,
    file_bytes: bytes,
    create_accounts: bool,
    update_existing: bool = False,
) -> ImportResult:
    """教师导入。

    `update_existing` 为真时,档案里已存在的教师(姓名＋身分末四码相同)改为更新而非报错——
    每学年职务异动要改几十位老师,逐笔手改不切实际(用户回报 #18)。
    更新时**只覆盖有填的栏位**,空白保留原值;要清掉行政职称请填「无」或「-」,
    数字栏要归零请填 0。既有教师的登录帐号一律不在这里处理(避免误建),
    请改用教师编辑界面的「建立登录帐号」。
    """
    result = ImportResult()
    subjects = {
        s.name: s
        for s in db.scalars(select(Subject).where(Subject.semester_id == semester_id))
    }
    existing = {
        (t.name, t.id_last4 or ""): t
        for t in db.scalars(select(Teacher).where(Teacher.semester_id == semester_id))
    }
    existing_usernames = set(db.scalars(select(User.username)))

    seen_keys: set[tuple[str, str]] = set()
    seen_usernames: set[str] = set()
    pending: list[tuple[Teacher, str | None]] = []   # 新增:(teacher, username or None)
    updates: list[tuple[Teacher, dict, list[Subject] | None]] = []  # 更新:(教师, 栏位, 科目)

    for idx, row in _data_rows(file_bytes):
        name = _cell(row, 0)
        if not name:
            result.errors.append(f"第 {idx} 列:姓名必填")
            continue
        id_last4 = _cell(row, 1)
        key = (name, id_last4 or "")
        current = existing.get(key)
        if key in seen_keys or (current is not None and not update_existing):
            result.errors.append(f"第 {idx} 列:教师「{name}」(末四码 {id_last4 or '无'})重复")
            continue
        seen_keys.add(key)

        subject_objs: list[Subject] = []
        subj_field = _cell(row, 2)
        subj_error = False
        if subj_field:
            names = [s.strip() for s in subj_field.replace(",", "、").split("、") if s.strip()]
            for sname in names:
                if sname not in subjects:
                    result.errors.append(f"第 {idx} 列:科目「{sname}」不存在")
                    subj_error = True
                    break
                subject_objs.append(subjects[sname])
        if subj_error:
            continue
        try:
            base_periods_cell = _parse_int(_cell(row, 3))
            admin_reduction_cell = _parse_int(_cell(row, 5))
        except ValueError as e:
            result.errors.append(f"第 {idx} 列:{e}")
            continue
        external_cell = _cell(row, 6)

        username = _cell(row, 7)
        if create_accounts and username and current is None:
            if username in existing_usernames or username in seen_usernames:
                result.errors.append(f"第 {idx} 列:登录帐号「{username}」重复")
                continue
            seen_usernames.add(username)

        email = _cell(row, 8)
        if email and not is_valid_email(email):
            result.errors.append(f"第 {idx} 列:Email「{email}」格式不正确")
            continue

        if current is not None:
            # 更新:只动有填的栏位。行政职称填「无」/「-」视为卸任(清空)
            fields: dict = {}
            title = _cell(row, 4)
            if title is not None:
                fields["admin_title"] = "" if title in _CLEAR_TOKENS else title
            if base_periods_cell is not None:
                fields["base_periods"] = base_periods_cell
            if admin_reduction_cell is not None:
                fields["admin_reduction"] = admin_reduction_cell
            if external_cell is not None:
                fields["is_external"] = external_cell == "是"
            for attr, value in (("email", email), ("phone", _cell(row, 9))):
                if value is not None:
                    fields[attr] = value
            updates.append((current, fields, subject_objs if subj_field else None))
            continue

        teacher = Teacher(
            semester_id=semester_id, name=name, id_last4=id_last4,
            base_periods=base_periods_cell or 0, admin_title=_cell(row, 4),
            admin_reduction=admin_reduction_cell or 0,
            is_external=(external_cell or "否") == "是",
            email=email, phone=_cell(row, 9),
            subjects=subject_objs,
        )
        pending.append((teacher, username if create_accounts else None))

    if result.errors:
        return result

    for teacher, username in pending:
        db.add(teacher)
        if username:
            # 绑定新建帐号:teacher.user 赋值于 flush 时写入 teachers.user_id
            teacher.user = User(
                username=username,
                password_hash=hash_password(settings.default_import_password),
                display_name=teacher.name,
                must_change_password=True,
                roles=[UserRole(role=Role.teacher.value)],
            )
    for teacher, fields, new_subjects in updates:
        for attr, value in fields.items():
            setattr(teacher, attr, value)
        if new_subjects is not None:
            teacher.subjects = new_subjects
    db.commit()
    result.imported = len(pending)
    result.updated = len(updates)
    return result


def _import_assignments(db: Session, semester_id: int, file_bytes: bytes) -> ImportResult:
    """单班配课导入(班级×科目×教师×周节数,可含一组连堂)。跑班群组于界面建立。"""
    result = ImportResult()
    classes = {
        c.name: c
        for c in db.scalars(select(ClassUnit).where(ClassUnit.semester_id == semester_id))
    }
    subjects = {
        s.name: s
        for s in db.scalars(select(Subject).where(Subject.semester_id == semester_id))
    }
    teachers = {
        t.name: t
        for t in db.scalars(select(Teacher).where(Teacher.semester_id == semester_id))
    }
    # (class_unit, subject, teacher_objs, periods, block, room_type)
    pending: list[tuple] = []
    for idx, row in _data_rows(file_bytes):
        class_name = _cell(row, 0)
        subj_name = _cell(row, 1)
        teacher_field = _cell(row, 2)
        if not class_name or not subj_name or not teacher_field:
            result.errors.append(f"第 {idx} 列:班级、科目、教师皆必填")
            continue
        if class_name not in classes:
            result.errors.append(f"第 {idx} 列:班级「{class_name}」不存在")
            continue
        if subj_name not in subjects:
            result.errors.append(f"第 {idx} 列:科目「{subj_name}」不存在")
            continue
        names = [s.strip() for s in teacher_field.replace(",", "、").split("、") if s.strip()]
        teacher_objs = []
        t_error = False
        for tname in names:
            if tname not in teachers:
                result.errors.append(f"第 {idx} 列:教师「{tname}」不存在")
                t_error = True
                break
            teacher_objs.append(teachers[tname])
        if t_error:
            continue
        try:
            periods = _parse_int(_cell(row, 3))
        except ValueError as e:
            result.errors.append(f"第 {idx} 列:每周节数 {e}")
            continue
        if not periods or periods < 1:
            result.errors.append(f"第 {idx} 列:每周节数必填且需大于 0")
            continue
        # 连堂(选填,长度与次数需成对)
        try:
            block_size = _parse_int(_cell(row, 4))
            block_count = _parse_int(_cell(row, 5))
        except ValueError as e:
            result.errors.append(f"第 {idx} 列:连堂 {e}")
            continue
        block: tuple[int, int] | None = None
        if block_size is not None or block_count is not None:
            if block_size is None or block_count is None:
                result.errors.append(f"第 {idx} 列:连堂长度与连堂次数需成对填写")
                continue
            if not 2 <= block_size <= 4 or block_count < 1:
                result.errors.append(f"第 {idx} 列:连堂长度需为 2-4、次数需大于 0")
                continue
            if block_size * block_count > periods:
                result.errors.append(f"第 {idx} 列:连堂总节数超过每周节数")
                continue
            block = (block_size, block_count)
        room_label = _cell(row, 6)
        room_type = None
        if room_label:
            if room_label not in ROOM_TYPE_BY_LABEL:
                result.errors.append(f"第 {idx} 列:场地类型「{room_label}」无效")
                continue
            room_type = ROOM_TYPE_BY_LABEL[room_label].value
        pending.append((classes[class_name], subjects[subj_name], teacher_objs, periods, block,
                        room_type))

    if result.errors:
        return result

    for class_unit, subject, teacher_objs, periods, block, room_type in pending:
        unit = get_or_create_single_unit(db, class_unit)
        assignment = CourseAssignment(
            semester_id=semester_id, scheduling_unit_id=unit.id, subject_id=subject.id,
            periods_per_week=periods, required_room_type=room_type,
        )
        db.add(assignment)
        db.flush()
        for i, t in enumerate(teacher_objs):
            assignment.teachers.append(
                AssignmentTeacher(teacher_id=t.id, is_lead=(i == 0))
            )
        if block is not None:
            assignment.block_rules.append(
                BlockRule(block_size=block[0], count_per_week=block[1])
            )

    # 超课时上限:整批一起检核。导入是「一次几百笔」的操作,
    # 逐列挡下只会让用户改一次汇一次;这里一次把所有超标的人列出来。
    db.flush()
    touched = {t.id for _, _, teacher_objs, _, _, _ in pending for t in teacher_objs}
    try:
        assert_within_overtime_limit(db, semester_id, touched)
    except DomainError as exc:
        db.rollback()
        result.errors.append(exc.message)
        result.imported = 0
        return result

    db.commit()
    result.imported = len(pending)
    return result


def run_import(
    db: Session,
    entity: str,
    semester_id: int,
    file_bytes: bytes,
    create_accounts: bool = False,
    update_existing: bool = False,
) -> ImportResult:
    if entity == "subjects":
        return _import_subjects(db, semester_id, file_bytes)
    if entity == "classes":
        return _import_classes(db, semester_id, file_bytes)
    if entity == "teachers":
        return _import_teachers(db, semester_id, file_bytes, create_accounts, update_existing)
    if entity == "assignments":
        return _import_assignments(db, semester_id, file_bytes)
    raise ValueError(f"未知的导入类型:{entity}")
