"""代课课时月结统计(M4-5,architecture.md §5.4)。

回答「这个月每位老师代了几节、其中几节要计课时费」。真相仍是 `substitution` 列
(处置决定)+ `affected_period`(受影响节次快照),这里依教师汇总。

**几个数字**:
- 代课节数:该教师接手的所有处置(代课/调课/并班),即他实际处理了几节。
- 计费节数(课时):其中 `counts_toward_hours` 为真、且不是日薪者。并班/自习默认不计、
  代课默认计(可覆写)。
- 日薪天数:计费方式为「外聘代课（按日计酬）」的代课,按日计酬——算的是「有几天」不是「几节」,
  所以这些节次不进课时(#36)。同一位请假教师同一天若由两位以上外聘代课（按日计酬）分担,
  每位各算一天并标示出来,怎么给付由学校决定(各县市规定不同,系统不代为判断)。

**跨月假单自动拆月**:以每一个 `affected_period` 自己的日期分月,不是以假单分月。
王师请 1/30~2/2 的假,1 月的节次进 1 月报表、2 月的进 2 月——不必特别处理。

**销假的节次不计**:销假会把未完成的节次转为 `cancelled`(那堂课没上),故排除;
已完成(completed)的节次即使事后销假仍保留(课上过了,课时照算)。
"""

import io
from dataclasses import dataclass, field
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.leave import LEAVE_TYPE_CN, AffectedPeriod, AffectedStatus
from app.models.substitution import (
    DAILY_FUNDING,
    DAILY_FUNDING_LEGACY,
    SUBSTITUTION_TYPE_CN,
    Substitution,
    SubstitutionType,
)

_Date = date


@dataclass(frozen=True, slots=True)
class StatDetail:
    """一列代课明细:某教师某节的处置。"""

    handler_teacher_id: int
    handler_name: str
    date: _Date
    period_no: int
    period_name: str
    class_names: str
    subject_name: str
    absent_teacher_name: str
    leave_type: str
    leave_type_label: str
    sub_type: str
    sub_type_label: str
    counts_toward_hours: bool
    funding_source: str
    pay_kind: str = "none"       # hourly 课时 / daily 日薪 / none 不计
    daily_shared: bool = False   # 日薪:这一天同一位请假教师的课由两位以上外聘代课（按日计酬）分担


PAY_KIND_CN = {"hourly": "课时", "daily": "日薪", "none": "不计"}


def pay_kind(sub: Substitution) -> str:
    if sub.type == SubstitutionType.substitute.value and sub.funding_source in (
        DAILY_FUNDING, DAILY_FUNDING_LEGACY,
    ):
        return "daily"
    return "hourly" if sub.counts_toward_hours else "none"


@dataclass
class TeacherSummary:
    teacher_id: int
    teacher_name: str
    handled_count: int = 0   # 代课节数(所有接手处置)
    billable_count: int = 0  # 课时计费节数(counts_toward_hours 为真,不含日薪)
    daily_days: int = 0      # 日薪天数
    daily_periods: int = 0   # 日薪那几天共代了几节(供核对)
    daily_shared_days: int = 0  # 其中几天是与其他外聘代课（按日计酬）老师分担的


@dataclass
class MonthlyReport:
    year: int
    month: int
    summaries: list[TeacherSummary] = field(default_factory=list)
    details: list[StatDetail] = field(default_factory=list)


def _next_month(year: int, month: int) -> date:
    return date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)


def monthly_report(
    db: Session,
    semester_id: int,
    year: int,
    month: int,
    *,
    teacher_id: int | None = None,
) -> MonthlyReport:
    """某月的代课课时统计。指定 teacher_id 则只统计该教师(教师个人查询用)。"""
    month_start = date(year, month, 1)
    month_end = _next_month(year, month)

    stmt = (
        select(Substitution, AffectedPeriod)
        .join(AffectedPeriod, Substitution.affected_period_id == AffectedPeriod.id)
        .where(
            Substitution.semester_id == semester_id,
            Substitution.handler_teacher_id.isnot(None),
            AffectedPeriod.status != AffectedStatus.cancelled.value,
            AffectedPeriod.date >= month_start,
            AffectedPeriod.date < month_end,
        )
    )
    if teacher_id is not None:
        stmt = stmt.where(Substitution.handler_teacher_id == teacher_id)

    report = MonthlyReport(year=year, month=month)
    summaries: dict[int, TeacherSummary] = {}
    rows = db.execute(stmt).all()
    shared = _shared_daily_days(db, semester_id, month_start, month_end)
    daily_days: dict[int, set[date]] = {}

    for sub, ap in rows:
        handler = sub.handler
        if handler is None:  # handler 已被移除(SET NULL 尚未反映在关联)
            continue
        leave = ap.leave_request
        kind = pay_kind(sub)
        is_shared = kind == "daily" and (leave.teacher_id, ap.date) in shared
        report.details.append(StatDetail(
            handler_teacher_id=handler.id,
            handler_name=handler.name,
            date=ap.date,
            period_no=ap.period_no,
            period_name=ap.period_name,
            class_names=ap.class_names,
            subject_name=ap.subject_name,
            absent_teacher_name=leave.teacher.name if leave.teacher else "(已移除)",
            leave_type=leave.leave_type,
            leave_type_label=LEAVE_TYPE_CN.get(leave.leave_type, leave.leave_type),
            sub_type=sub.type,
            sub_type_label=SUBSTITUTION_TYPE_CN.get(sub.type, sub.type),
            counts_toward_hours=sub.counts_toward_hours,
            funding_source=sub.funding_source,
            pay_kind=kind,
            daily_shared=is_shared,
        ))
        s = summaries.get(handler.id)
        if s is None:
            s = TeacherSummary(teacher_id=handler.id, teacher_name=handler.name)
            summaries[handler.id] = s
        s.handled_count += 1
        if kind == "hourly":
            s.billable_count += 1
        elif kind == "daily":
            s.daily_periods += 1
            days = daily_days.setdefault(handler.id, set())
            if ap.date not in days:
                days.add(ap.date)
                s.daily_days += 1
                if is_shared:
                    s.daily_shared_days += 1

    report.details.sort(key=lambda d: (d.handler_name, d.date, d.period_no))
    report.summaries = sorted(summaries.values(), key=lambda s: s.teacher_name)
    return report


def _shared_daily_days(
    db: Session, semester_id: int, month_start: date, month_end: date
) -> set[tuple[int, date]]:
    """(请假教师, 日期):那一天他的课由两位以上「外聘代课（按日计酬）」老师分担。

    一律看全校,不受 teacher_id 筛选影响——教师个人查询时也要看得到「这天是分担的」。
    """
    handlers: dict[tuple[int, date], set[int]] = {}
    for sub, ap in db.execute(
        select(Substitution, AffectedPeriod)
        .join(AffectedPeriod, Substitution.affected_period_id == AffectedPeriod.id)
        .where(
            Substitution.semester_id == semester_id,
            Substitution.handler_teacher_id.isnot(None),
            Substitution.type == SubstitutionType.substitute.value,
            Substitution.funding_source.in_([DAILY_FUNDING, DAILY_FUNDING_LEGACY]),
            AffectedPeriod.status != AffectedStatus.cancelled.value,
            AffectedPeriod.date >= month_start,
            AffectedPeriod.date < month_end,
        )
    ).all():
        if sub.handler_teacher_id is None:
            continue
        key = (ap.leave_request.teacher_id, ap.date)
        handlers.setdefault(key, set()).add(sub.handler_teacher_id)
    return {key for key, ids in handlers.items() if len(ids) > 1}


_DETAIL_HEADERS = (
    "教师", "日期", "节次", "班级", "科目", "原任教师", "假别", "处置", "计费", "经费来源",
    "计酬", "备注",
)
_SUMMARY_HEADERS = ("教师", "代课节数", "课时计费节数", "日薪天数", "日薪节数", "备注")
_SHARED_NOTE = "当日与其他按日代课分担"


def build_workbook(report: MonthlyReport) -> bytes:
    """导出两张表:汇总(每位教师)+ 明细(逐节)。回传 xlsx bytes。"""
    wb = Workbook()

    ws_sum = wb.active
    ws_sum.title = "汇总"
    ws_sum.append(list(_SUMMARY_HEADERS))
    for s in report.summaries:
        note = f"其中 {s.daily_shared_days} 天与其他外聘代课（按日计酬）分担" if s.daily_shared_days else ""
        ws_sum.append([
            s.teacher_name, s.handled_count, s.billable_count,
            s.daily_days, s.daily_periods, note,
        ])

    ws_detail = wb.create_sheet("明细")
    ws_detail.append(list(_DETAIL_HEADERS))
    for d in report.details:
        ws_detail.append([
            d.handler_name, d.date.isoformat(), d.period_name, d.class_names, d.subject_name,
            d.absent_teacher_name, d.leave_type_label, d.sub_type_label,
            "是" if d.counts_toward_hours else "否", d.funding_source,
            PAY_KIND_CN[d.pay_kind], _SHARED_NOTE if d.daily_shared else "",
        ])

    for ws in (ws_sum, ws_detail):
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 14

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
