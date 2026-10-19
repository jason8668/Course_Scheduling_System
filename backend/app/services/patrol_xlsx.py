"""巡堂表导出 Excel(v1.2.14)。

打印页印出来就不能改了;学校实务上常要在发出去之前再调整——加栏位、填巡堂人员或
值日、并进自己既有的巡堂纪录档(用户建议)。这里把同一份资料(`patrol_sheets.build`)
写成可编辑的活页簿,版面比照打印页:

- 每张巡堂纪录一个分页(每天的上午、下午;班级多时再分页),班级为栏、节次为列,
  每节五列:科目/教师、教室、授课情形、学生学习、备注;最右边是巡堂签名栏。
- 每张分组巡堂单(社团等)一个分页。
- 已设成 A4 横向、缩成一页宽,下载后不改也能直接印。

定位是「拿来改」而不是取代打印页:字型与格线不会和打印页一模一样。
导出之后若又有新的调代课,档案不会跟著变,要重新导出。
"""

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.patrol_sheets import GroupList, PatrolPage, PatrolSheets

_WEEKDAYS = "一二三四五六日"
_ROW_LABELS = ("科目\n教师", "教室", "授课情形", "学生学习", "备注")

_THIN = Side(style="thin", color="000000")
_MEDIUM = Side(style="medium", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _date_label(day: date) -> str:
    """「115年10月13日 星期二」:纸本用民国年。"""
    return f"{day.year - 1911}年{day.month}月{day.day}日　星期{_WEEKDAYS[day.isoweekday() - 1]}"


def _range(first: int, last: int) -> str:
    return f"第 {first} 节" if first == last else f"第 {first}~{last} 节"


def _sheet_name(base: str, used: set[str]) -> str:
    # Excel 分页名 ≤31 字、不可含 : \ / ? * [ ];重复时加流水号
    clean = base
    for ch in ":\\/?*[]":
        clean = clean.replace(ch, " ")
    clean = clean[:28].strip() or "巡堂表"
    name, i = clean, 2
    while name in used:
        name = f"{clean[:25]}({i})"
        i += 1
    used.add(name)
    return name


def _box(ws: Worksheet, first_row: int, last_row: int, last_col: int) -> None:
    """整块画细格线,并把这一节的上缘加粗(纸本每一节之间是粗线)。"""
    for row in ws.iter_rows(min_row=first_row, max_row=last_row, min_col=1, max_col=last_col):
        for cell in row:
            cell.border = Border(
                left=_THIN, right=_THIN, bottom=_THIN,
                top=_MEDIUM if cell.row == first_row else _THIN,
            )


def _landscape(ws: Worksheet) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.35
    ws.page_margins.top = ws.page_margins.bottom = 0.4


def _write_page(ws: Worksheet, title: str, legend: str, page: PatrolPage) -> None:
    n = len(page.classes)
    last_col = n + 3                      # 节次 | 列名 | 各班… | 巡堂签名
    sign_col = last_col

    periods = _range(page.first_ordinal, page.last_ordinal)
    heading = f"{title}巡堂纪录　{_date_label(page.date)}　{periods}"
    if page.table_name:
        heading += f"　({page.table_name})"
    ws.cell(row=1, column=1, value=heading).font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.row_dimensions[1].height = 26

    ws.cell(row=2, column=1, value="节次")
    for k, name in enumerate(page.classes):
        ws.cell(row=2, column=3 + k, value=name)
    ws.cell(row=2, column=sign_col, value="巡堂\n签名")
    for cell in ws[2]:
        cell.alignment = _CENTER
        cell.font = Font(bold=True)
        cell.border = _BORDER
    ws.row_dimensions[2].height = 30

    r = 3
    for row in page.rows:
        first = r
        for offset, label in enumerate(_ROW_LABELS):
            ws.cell(row=r + offset, column=2, value=label)
        for k, cell in enumerate(row.cells):
            col = 3 + k
            lesson = "\n".join(x for x in (cell.subject, cell.teacher) if x)
            ws.cell(row=r, column=col, value=lesson or None)
            ws.cell(row=r + 1, column=col, value=cell.room or None)
            ws.cell(row=r + 4, column=col, value=cell.note or None)
        ws.cell(row=first, column=1, value=row.ordinal)
        ws.merge_cells(start_row=first, start_column=1, end_row=first + 4, end_column=1)
        ws.merge_cells(
            start_row=first, start_column=sign_col, end_row=first + 4, end_column=sign_col)
        _box(ws, first, first + 4, last_col)
        for offset in range(5):
            ws.row_dimensions[r + offset].height = 34 if offset == 0 else 20
            for c in range(1, last_col + 1):
                ws.cell(row=r + offset, column=c).alignment = _CENTER
        r += 5

    # 记录说明 + 注记事项:各占一半宽度
    lines = [line for line in legend.split("\n") if line]
    half = max(2, (last_col + 1) // 2)
    ws.cell(row=r, column=1, value="记录\n说明")
    ws.cell(row=r, column=2, value="\n".join(lines))
    ws.cell(row=r, column=half + 1, value="注记事项")
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=half)
    if half + 2 <= last_col:
        ws.merge_cells(start_row=r, start_column=half + 2, end_row=r, end_column=last_col)
    for c in range(1, last_col + 1):
        cell = ws.cell(row=r, column=c)
        cell.border = Border(left=_THIN, right=_THIN, bottom=_THIN, top=_MEDIUM)
        cell.alignment = _LEFT if c == 2 else _CENTER
    ws.row_dimensions[r].height = max(48, 17 * len(lines))

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 10
    for k in range(n):
        ws.column_dimensions[get_column_letter(3 + k)].width = 13
    ws.column_dimensions[get_column_letter(sign_col)].width = 8
    ws.freeze_panes = "C3"
    _landscape(ws)


def _write_list(ws: Worksheet, title: str, lst: GroupList) -> None:
    ws.cell(row=1, column=1, value=title)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
    heading = f"{lst.group_name}巡堂单　{_date_label(lst.date)}　{lst.period_name}"
    ws.cell(row=2, column=1, value=heading).font = Font(bold=True, size=14)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=6)
    ws.row_dimensions[2].height = 26

    for c, name in enumerate(("编号", "名称", "教师", "地点", "上课状况", "学生表现"), start=1):
        cell = ws.cell(row=3, column=c, value=name)
        cell.font = Font(bold=True)
        cell.alignment = _CENTER
        cell.border = _BORDER
    for i, item in enumerate(lst.items):
        r = 4 + i
        teacher = f"{item.teacher}\n({item.note})" if item.note else item.teacher
        for c, value in enumerate((item.no, item.subject, teacher, item.room, None, None), start=1):
            cell = ws.cell(row=r, column=c, value=value or None)
            cell.alignment = _CENTER
            cell.border = _BORDER
        ws.row_dimensions[r].height = 30
    for letter, width in zip("ABCDEF", (7, 22, 18, 20, 16, 16), strict=True):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A4"
    _landscape(ws)
    ws.page_setup.orientation = "portrait"      # 清单一份直印刚好一页


def to_xlsx(sheets: PatrolSheets) -> bytes:
    """巡堂纪录与分组巡堂单各自成页的活页簿。分页依日期排,同一天先主表、后分组单。"""
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    days = sorted({p.date for p in sheets.pages} | {g.date for g in sheets.group_lists})
    for day in days:
        for page in (p for p in sheets.pages if p.date == day):
            half = "上午" if page.first_ordinal == 1 else "下午"
            table = f" {page.table_name}" if page.table_name else ""
            ws = wb.create_sheet(_sheet_name(f"{day:%m-%d} {half}{table}", used))
            _write_page(ws, sheets.title, sheets.legend, page)
        for lst in (g for g in sheets.group_lists if g.date == day):
            name = f"{day:%m-%d} {lst.group_name} {lst.period_name}"
            ws = wb.create_sheet(_sheet_name(name, used))
            _write_list(ws, sheets.title, lst)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
