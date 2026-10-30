"""课表导出:把已发布课表摊成格线,产生 Excel / PDF(HTML)/PNG(M5-1)。

三种对象(班级 / 教师 / 场地)共用同一个格线模型 `Grid`,再交给各格式的渲染器,
确保三种格式与界面课表内容一致(验收①)。资料来源一律是**已发布**课表(D4 快照)。

- Excel 由 openpyxl 产生,可在 api 同步跑(轻量)。
- PDF 需 WeasyPrint(系统依赖+中文字型只在 worker),故 PDF/PNG 走 worker 背景渲染;
  本模组只负责产出 HTML,实际 write_pdf 在 `app.workers.export_job`。
"""

import io
import zipfile
from dataclasses import dataclass, field
from datetime import date

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.week_parity import PARITY_SHORT
from app.models.basedata import ClassUnit, Room, Teacher
from app.models.period import Period, PeriodTable, PeriodType
from app.models.semester import Semester
from app.models.timetable import ScheduleEntry, Timetable, TimetableStatus
from app.services import period_tables as pt_service
from app.services import settings as app_settings

_WEEKDAYS = ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]


class ExportError(Exception):
    """导出前置条件不成立(呼叫端转为 4xx)。"""


# ── 格线模型 ────────────────────────────────────────────────
@dataclass(frozen=True, slots=True)
class Cell:
    lines: tuple[str, ...] = ()   # 科目 / 教师或班级 / 教室(逐行)
    span: int = 1                 # 连堂占几列
    covered: bool = False         # 被上方连堂覆盖,不绘


@dataclass
class Row:
    period_no: int
    label: str
    is_regular: bool
    cells: list[Cell] = field(default_factory=list)  # 长度 = num_weekdays


@dataclass
class Grid:
    title: str
    num_weekdays: int
    rows: list[Row] = field(default_factory=list)
    # 行政值日行:按周一..周日顺序的教师名,空 tuple = 本学期没排轮值,不画这行。
    # 与课表打印页的值日行同一口径(每周轮值),导出文件与打印页内容一致。
    duty: tuple[str, ...] = ()

    @property
    def weekday_names(self) -> list[str]:
        return _WEEKDAYS[: self.num_weekdays]


@dataclass(frozen=True, slots=True)
class Meta:
    school_name: str
    semester_label: str
    timetable_name: str
    printed_on: date


@dataclass(frozen=True, slots=True)
class _EntryView:
    weekday: int
    period_no: int
    span: int
    subject: str
    teachers: str
    classes: str
    room: str
    teacher_ids: frozenset[int]
    class_ids: frozenset[int]
    room_id: int | None
    table_id: int | None = None   # 这一格所属的节次表(取第一个班级的)
    week_parity: str | None = None  # 单双周(v1.3.0):odd / even;None = 每周


class _Published:
    """一份已发布课表的导出来源:格位、节次表、对象清单一次捞齐。"""

    def __init__(self, db: Session, semester_id: int) -> None:
        self.db = db
        semester = db.get(Semester, semester_id)
        if semester is None:
            raise ExportError("找不到学期")
        timetable = db.scalar(
            select(Timetable).where(
                Timetable.semester_id == semester_id,
                Timetable.status == TimetableStatus.published.value,
            )
        )
        if timetable is None:
            raise ExportError("此学期尚无已发布的课表")
        self.semester: Semester = semester
        self.timetable: Timetable = timetable
        self.entries = [
            _view(e)
            for e in db.scalars(
                select(ScheduleEntry).where(ScheduleEntry.timetable_id == self.timetable.id)
            )
        ]
        self.tables = {
            t.id: t for t in db.scalars(
                select(PeriodTable).where(PeriodTable.semester_id == semester_id)
            )
        }
        # 行政值日每周轮值(周一=1..周日=7),供导出值日行;没排轮值的学期是空表。
        from app.models.duty import DutyPattern
        patterns = db.scalars(
            select(DutyPattern).where(DutyPattern.semester_id == semester_id)
        ).all()
        teacher_names = {
            t.id: t.name for t in db.scalars(
                select(Teacher).where(Teacher.semester_id == semester_id)
            )
        }
        by_weekday = {p.weekday: teacher_names.get(p.teacher_id, "") for p in patterns}
        self.duty = tuple(by_weekday.get(wd, "") for wd in range(1, 8))

    def meta(self) -> Meta:
        from app.core import clock
        return Meta(
            school_name=app_settings.school_name(self.db),
            semester_label=self.semester.label,
            timetable_name=self.timetable.name,
            printed_on=clock.school_today(),
        )

    def default_table(self) -> PeriodTable | None:
        tables = list(self.tables.values())
        return next((t for t in tables if t.is_default), tables[0] if tables else None)

    def class_table(self, cls: ClassUnit) -> PeriodTable | None:
        t = pt_service.resolve_period_table(self.db, cls)
        return t or self.default_table()


def _view(e: ScheduleEntry) -> _EntryView:
    a = e.assignment
    su = a.scheduling_unit
    room = e.room if e.room is not None else a.room
    return _EntryView(
        weekday=e.weekday, period_no=e.period_no, span=e.span,
        subject=a.subject.name,
        teachers="、".join(at.teacher.name for at in a.teachers),
        classes="、".join(m.class_unit.name for m in su.members),
        room=room.name if room else "",
        teacher_ids=frozenset(at.teacher_id for at in a.teachers),
        class_ids=frozenset(m.class_unit_id for m in su.members),
        room_id=e.effective_room_id,
        # 没指定节次表的班级是 None(=学期默认表,也就是教师课表格线用的那套)
        table_id=su.members[0].class_unit.period_table_id if su.members else None,
        week_parity=a.week_parity,
    )


def _hhmm(p: Period) -> str:
    return f"{p.start_time:%H:%M}–{p.end_time:%H:%M}" if p.start_time and p.end_time else ""


def _time_note(e: _EntryView, shown: PeriodTable, tables: dict[int, PeriodTable]) -> str:
    """这一格的实际上课时间——只有它属于另一套节次表、且时间与所在列不同时才回传。

    教师/场地课表只能照一套节次表画格线。初中小、完全中学的跨部教师,
    另一部的课排在「同节次号」那一列,但实际时间不一样(小学 40 分、初中 45 分);
    不标出来的话,课表上的时间就是错的。
    """
    own = tables.get(e.table_id) if e.table_id is not None else None
    if own is None or own.id == shown.id:
        return ""

    def at(t: PeriodTable) -> Period | None:
        return next(
            (p for p in t.periods if p.weekday == e.weekday and p.period_no == e.period_no), None)
    mine, row = at(own), at(shown)
    if mine is None or not _hhmm(mine):
        return ""
    if row is not None and (row.start_time, row.end_time) == (mine.start_time, mine.end_time):
        return ""
    return _hhmm(mine)


def _grid_from(
    table: PeriodTable | None,
    title: str,
    entries: list[_EntryView],
    lines_of,
    tables: dict[int, PeriodTable] | None = None,
    duty: tuple[str, ...] = (),
) -> Grid:
    """把某对象的格位排进节次表格线。`lines_of(entry)` 决定每格显示哪几行。

    给了 `tables`(教师/场地课表)就会在跨节次表的格子多一行实际时间。
    """
    if table is None:
        return Grid(title=title, num_weekdays=5, rows=[])
    num_weekdays = table.num_weekdays
    periods = sorted(table.periods, key=lambda p: (p.period_no, p.weekday))
    # 每个 period_no 取一个代表(名称/类型),weekday 小者优先
    by_no: dict[int, Period] = {}
    for p in periods:
        by_no.setdefault(p.period_no, p)
    # 另一套节次表才有的节次(初中七节、高中八节的完全中学,跨部老师的第八节):
    # 这套格线没有那一列,不补的话那堂课会整个从课表上消失。只补用得到的列。
    for e in entries:
        own = (tables or {}).get(e.table_id) if e.table_id is not None else None
        for k in range(e.span):
            pno = e.period_no + k
            if pno in by_no or own is None:
                continue
            extra = next((p for p in sorted(own.periods, key=lambda p: p.weekday)
                          if p.period_no == pno), None)
            if extra is not None:
                by_no[pno] = extra
    order = sorted(by_no)

    # (weekday, period_no) → 这一格的课。单双周的格子有单周、双周两堂(v1.3.0)
    placed: dict[tuple[int, int], list[_EntryView]] = {}
    for e in entries:
        placed.setdefault((e.weekday, e.period_no), []).append(e)
    covered: set[tuple[int, int]] = set()
    for e in entries:
        for k in range(1, e.span):
            covered.add((e.weekday, e.period_no + k))

    grid = Grid(title=title, num_weekdays=num_weekdays)
    # 值日行只在确实排了轮值时出现;没排的那天显示 —,与打印页一致。
    if any(duty):
        grid.duty = tuple(duty[wd - 1] if wd <= len(duty) and duty[wd - 1] else "—"
                          for wd in range(1, num_weekdays + 1))
    for pno in order:
        rep = by_no[pno]
        is_regular = rep.type == PeriodType.regular.value
        row = Row(period_no=pno, label=rep.name, is_regular=is_regular)
        for wd in range(1, num_weekdays + 1):
            if (wd, pno) in covered:
                row.cells.append(Cell(covered=True))
            elif (wd, pno) in placed:
                here = placed[(wd, pno)]
                alternating = sorted(
                    (x for x in here if x.week_parity), key=lambda x: x.week_parity != "odd")
                if alternating:
                    # 「单 生科 甲师」「双 资科 乙师」:一行一周,一眼看出隔周轮流
                    e = alternating[0]
                    lines = [
                        f"{PARITY_SHORT[x.week_parity or '']} " + " ".join(lines_of(x))
                        for x in alternating
                    ]
                else:
                    e = here[-1]
                    lines = list(lines_of(e))
                note = _time_note(e, table, tables) if tables else ""
                if note:
                    lines.append(note)
                row.cells.append(Cell(lines=tuple(lines), span=e.span))
            else:
                row.cells.append(Cell())
        grid.rows.append(row)
    return grid


# ── 三种对象 → Grid ─────────────────────────────────────────
def _class_lines(e: _EntryView) -> list[str]:
    return [x for x in (e.subject, e.teachers, e.room) if x]


def _teacher_lines(e: _EntryView) -> list[str]:
    return [x for x in (e.subject, e.classes, e.room) if x]


def _room_lines(e: _EntryView) -> list[str]:
    return [x for x in (e.subject, e.classes, e.teachers) if x]


def class_grid(pub: _Published, cls: ClassUnit) -> Grid:
    entries = [e for e in pub.entries if cls.id in e.class_ids]
    return _grid_from(
        pub.class_table(cls), f"{cls.grade}年{cls.name} 课表", entries,
        _class_lines, duty=pub.duty)


def teacher_grid(pub: _Published, teacher: Teacher) -> Grid:
    entries = [e for e in pub.entries if teacher.id in e.teacher_ids]
    return _grid_from(
        pub.default_table(), f"{teacher.name} 课表", entries,
        _teacher_lines, pub.tables, duty=pub.duty)


def room_grid(pub: _Published, room: Room) -> Grid:
    entries = [e for e in pub.entries if e.room_id == room.id]
    return _grid_from(
        pub.default_table(), f"{room.name} 课表", entries,
        _room_lines, pub.tables, duty=pub.duty)


def build_grid(db: Session, semester_id: int, view: str, target_id: int) -> tuple[Grid, Meta]:
    pub = _Published(db, semester_id)
    if view == "class":
        obj = db.get(ClassUnit, target_id)
        if obj is None or obj.semester_id != semester_id:
            raise ExportError("找不到班级")
        return class_grid(pub, obj), pub.meta()
    if view == "teacher":
        obj_t = db.get(Teacher, target_id)
        if obj_t is None or obj_t.semester_id != semester_id:
            raise ExportError("找不到教师")
        return teacher_grid(pub, obj_t), pub.meta()
    if view == "room":
        obj_r = db.get(Room, target_id)
        if obj_r is None or obj_r.semester_id != semester_id:
            raise ExportError("找不到场地")
        return room_grid(pub, obj_r), pub.meta()
    raise ExportError(f"未知的查看类型:{view}")


def _classes(db: Session, semester_id: int) -> list[ClassUnit]:
    return list(db.scalars(
        select(ClassUnit).where(ClassUnit.semester_id == semester_id)
        .order_by(ClassUnit.grade, ClassUnit.name)
    ))


def _teachers(db: Session, semester_id: int) -> list[Teacher]:
    """在职教师,依姓名排序。

    没有排到课的也列入(会是一张空白课表)——全校发放时宁可多一张,也不要漏人
    (用户回报 #22);离职教师则不列入,与代课推荐的口径一致。
    """
    return list(db.scalars(
        select(Teacher).where(
            Teacher.semester_id == semester_id, Teacher.is_active.is_(True)
        ).order_by(Teacher.name)
    ))


def school_workbook(db: Session, semester_id: int, view: str = "class") -> bytes:
    """全校总表:一个 Excel 档,每个班级(或每位教师)一个分页。"""
    pub = _Published(db, semester_id)
    if view == "teacher":
        grids = [teacher_grid(pub, t) for t in _teachers(db, semester_id)]
    else:
        grids = [class_grid(pub, c) for c in _classes(db, semester_id)]
    return grids_to_xlsx(grids, pub.meta())


def batch_zip(db: Session, semester_id: int, view: str = "class") -> bytes:
    """批次导出:每个班级(或每位教师)各一个 Excel 档,打包成 zip。"""
    pub = _Published(db, semester_id)
    meta = pub.meta()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        if view == "teacher":
            for t in _teachers(db, semester_id):
                z.writestr(f"{t.name}.xlsx", grids_to_xlsx([teacher_grid(pub, t)], meta))
        else:
            for c in _classes(db, semester_id):
                z.writestr(f"{c.grade}年{c.name}.xlsx",
                           grids_to_xlsx([class_grid(pub, c)], meta))
    return buf.getvalue()


def class_batch_zip(db: Session, semester_id: int) -> bytes:
    """批次导出全部班级(保留旧名称,供既有呼叫端使用)。"""
    return batch_zip(db, semester_id, "class")


# ── Excel 渲染 ──────────────────────────────────────────────
def _safe_sheet_title(title: str, used: set[str]) -> str:
    # Excel 分页名 ≤31 字、不可含 : \ / ? * [ ]
    clean = title
    for ch in ':\\/?*[]':
        clean = clean.replace(ch, " ")
    clean = clean[:28].strip() or "课表"
    name, i = clean, 1
    while name in used:
        name = f"{clean[:26]}~{i}"
        i += 1
    used.add(name)
    return name


def grids_to_xlsx(grids: list[Grid], meta: Meta) -> bytes:
    """每张 Grid 一个分页。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    head_fill = PatternFill("solid", fgColor="E8E8E8")
    shade = PatternFill("solid", fgColor="F5F5F5")
    used: set[str] = set()

    for grid in grids:
        ws = wb.create_sheet(_safe_sheet_title(grid.title, used))
        ws.append([grid.title])
        ws.append([f"{meta.school_name}　{meta.semester_label}　打印日:{meta.printed_on}"])
        ws["A1"].font = Font(bold=True, size=14)
        header = ["节次", *grid.weekday_names]
        ws.append(header)
        head_row = ws.max_row
        for c in ws[head_row]:
            c.font = Font(bold=True)
            c.fill = head_fill
            c.alignment = center
            c.border = border

        for row in grid.rows:
            values = [row.label]
            for cell in row.cells:
                values.append("\n".join(cell.lines) if cell.lines else "")
            ws.append(values)
            r = ws.max_row
            for ci, c in enumerate(ws[r]):
                c.alignment = center
                c.border = border
                if ci == 0 or not row.is_regular:
                    c.fill = shade
            # 连堂:垂直合并
            for ci, cell in enumerate(row.cells):
                if cell.span > 1:
                    col = ci + 2
                    ws.merge_cells(start_row=r, start_column=col,
                                   end_row=r + cell.span - 1, end_column=col)

        # 行政值日行:与打印页同一行,没排轮值的学期 grid.duty 为空,不画。
        if grid.duty:
            ws.append(["行政值日", *grid.duty])
            r = ws.max_row
            for ci, c in enumerate(ws[r]):
                c.alignment = center
                c.border = border
                if ci == 0:
                    c.font = Font(bold=True)
                    c.fill = shade

        ws.column_dimensions["A"].width = 10
        for col in range(2, 2 + grid.num_weekdays):
            ws.column_dimensions[ws.cell(row=head_row, column=col).column_letter].width = 16
        ws.freeze_panes = ws.cell(row=head_row + 1, column=2)

    if not wb.sheetnames:  # 全空:给一张空白页避免坏档
        wb.create_sheet("课表")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── HTML 渲染(供 worker 转 PDF/PNG)─────────────────────────
def _esc(s: str) -> str:
    return (s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def grid_to_html(grid: Grid, meta: Meta) -> str:
    """单一对象的 A4 直式课表 HTML。中文由 worker 映像内嵌的 Noto CJK 呈现。"""
    head_cells = "".join(f"<th>{_esc(w)}</th>" for w in grid.weekday_names)
    subtitle = (f"{_esc(meta.school_name)}　{_esc(meta.semester_label)}　"
                f"{_esc(meta.timetable_name)}　打印日:{meta.printed_on}")
    body_rows = []
    for row in grid.rows:
        tds = [f'<th class="pno">{_esc(row.label)}</th>']
        for cell in row.cells:
            if cell.covered:
                continue
            cls = "cell" if row.is_regular else "cell rest"
            span = f' rowspan="{cell.span}"' if cell.span > 1 else ""
            inner = "<br>".join(_esc(x) for x in cell.lines)
            tds.append(f'<td class="{cls}"{span}>{inner}</td>')
        body_rows.append(f"<tr>{''.join(tds)}</tr>")
    if grid.duty:
        duty_tds = "".join(f"<td class=\"cell\">{_esc(d)}</td>" for d in grid.duty)
        body_rows.append(f"<tr><th class=\"pno\">行政值日</th>{duty_tds}</tr>")
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>
@page {{ size: A4 portrait; margin: 14mm; }}
* {{ font-family: "Noto Sans CJK TC", "Noto Sans TC", sans-serif; }}
h1 {{ font-size: 18px; text-align: center; margin: 0 0 4px; }}
.meta {{ text-align: center; font-size: 12px; color: #444; margin-bottom: 10px; }}
table {{ border-collapse: collapse; width: 100%; table-layout: fixed; }}
th, td {{ border: 1px solid #333; padding: 4px; text-align: center; font-size: 12px;
  vertical-align: middle; word-break: break-all; }}
thead th {{ background: #e8e8e8; }}
.pno {{ background: #f2f2f2; width: 60px; font-size: 11px; }}
td.cell {{ height: 46px; }}
td.rest {{ background: #f7f7f7; color: #888; }}
</style></head><body>
<h1>{_esc(grid.title)}</h1>
<div class="meta">{subtitle}</div>
<table><thead><tr><th class="pno">节次</th>{head_cells}</tr></thead>
<tbody>{''.join(body_rows)}</tbody></table>
</body></html>"""
